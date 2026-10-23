"""Render a report (cover details + Markdown body) to a .docx matching docs/style-spec.md.

The Markdown is parsed by pandoc into its JSON AST, then every block is written with python-docx
so that fonts, spacing, numbering, borders and equations are fully under our control.
"""

from __future__ import annotations

import copy
import json
import tempfile
import re
from dataclasses import dataclass, field
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import ImageOps

from .omml import M, latex_to_omml
from .profile import Profile
from .styles import HeadingStyle, TemplateStyle

PAGE_W, PAGE_H = 8.5, 11.0
MARGIN_TOP, MARGIN_SIDE = 0.75, 0.79
CONTENT_W = PAGE_W - 2 * MARGIN_SIDE
MAX_IMG_H = 8.0
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}
_PAGEBREAK_RAW = re.compile(r"^\s*(<!--\s*page\s*break\s*-->|\\newpage|\\pagebreak|\\clearpage)\s*$", re.I)
_UNNUMBERED = re.compile(r"^(references|bibliography|appendix\b.*|acknowledge?ments?)$", re.I)
_QUESTION_PREFIX = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?\d+\s*[.):]\s*", re.I)


@dataclass
class CoverInfo:
    title: str
    module: str = ""
    submission_date: str = ""


@dataclass
class BuildResult:
    path: Path
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- low-level XML helpers


def _set_fonts(rpr_owner, name: str) -> None:
    """Force a font on every script slot and drop theme-font overrides."""
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _shade(ppr_or_tcpr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr_or_tcpr.append(shd)


def _page_borders(sect_pr) -> None:
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "28")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    # pgBorders must come after pgMar/paperSrc in CT_SectPr; insert right after pgMar.
    pg_mar = sect_pr.find(qn("w:pgMar"))
    pg_mar.addnext(borders)


def _para_box(p) -> None:
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "BFBFBF")
        bdr.append(el)
    ppr.append(bdr)
    _shade(ppr, "F5F5F5")


# --------------------------------------------------------------------------- numbering


class Numbering:
    """Creates list definitions directly in numbering.xml so each list can restart at its own value."""

    BULLETS = ("•", "o", "▪")
    ORDERED_FMT = {
        "Decimal": ("decimal", "%{n}."),
        "DefaultStyle": ("decimal", "%{n}."),
        "Example": ("decimal", "%{n}."),
        "LowerRoman": ("lowerRoman", "%{n}."),
        "UpperRoman": ("upperRoman", "%{n}."),
        "LowerAlpha": ("lowerLetter", "%{n})"),
        "UpperAlpha": ("upperLetter", "%{n}."),
    }

    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        self._next_abs = 100
        self._next_num = 100

    def _lvl(self, ilvl: int, fmt: str, text: str, left: int, hanging: int, start: int = 1, font: str | None = None):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", str(start)), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            lvl.append(el)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        ppr.append(ind)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), font)
            rpr.append(rf)
            lvl.append(rpr)
        return lvl

    def _register(self, levels) -> int:
        abs_id, num_id = self._next_abs, self._next_num
        self._next_abs += 1
        self._next_num += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        abstract.append(mlt)
        for lvl in levels:
            abstract.append(lvl)
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            self.root.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num.append(ref)
        self.root.append(num)
        return num_id

    def bullet(self, base_left: int = 720) -> int:
        return self._register(
            [self._lvl(i, "bullet", self.BULLETS[i % 3], base_left + 360 * i, 360, font="Arial") for i in range(6)]
        )

    def ordered(self, style: str = "Decimal", start: int = 1, base_left: int = 720) -> int:
        fmt0, text0 = self.ORDERED_FMT.get(style, ("decimal", "%{n}."))
        nested = [("lowerLetter", "%{n})"), ("lowerRoman", "%{n}."), ("decimal", "%{n}.")]
        levels = [self._lvl(0, fmt0, text0.format(n=1), base_left, 360, start)]
        for i in range(1, 6):
            fmt, text = nested[(i - 1) % 3]
            levels.append(self._lvl(i, fmt, text.format(n=i + 1), base_left + 360 * i, 360))
        return self._register(levels)

    def questions(self) -> int:
        return self._register([self._lvl(0, "decimal", "%1.", 360, 360)])


def _apply_num(p, num_id: int, ilvl: int) -> None:
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(lvl)
    numpr.append(nid)
    # numPr belongs near the start of pPr (after pStyle/keepNext etc.); python-docx tolerates insert at 0
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is not None:
        pstyle.addnext(numpr)
    else:
        ppr.insert(0, numpr)


# --------------------------------------------------------------------------- the builder


class ReportBuilder:
    def __init__(self, style: TemplateStyle, profile: Profile, assets_dir: Path | None = None):
        self.s = style
        self.profile = profile
        self.assets_dir = assets_dir
        self.warnings: list[str] = []
        self.doc = Document()
        self.num = Numbering(self.doc)
        self._question_num: int | None = None
        self._after_backmatter = False
        self._math: dict[int, object] = {}
        self._fig_n = 0
        self._tab_n = 0
        self._tmp = tempfile.TemporaryDirectory()
        self._setup_page()
        self._setup_styles()

    # ---- setup

    def _setup_page(self) -> None:
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(PAGE_W), Inches(PAGE_H)
        sec.top_margin = Inches(MARGIN_TOP)
        sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(MARGIN_SIDE)
        sec.header_distance = sec.footer_distance = Inches(0.5)
        _page_borders(sec._sectPr)

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = self.s.body_font
        normal.font.size = Pt(self.s.body_size)
        _set_fonts(normal.element, self.s.body_font)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(8), 1.15

        for level, hs in ((1, self.s.h1), (2, self.s.h2), (3, self.s.h3)):
            st = styles[f"Heading {level}"]
            self._style_heading(st, hs)

        for name in ("Title", "Caption"):
            st = styles[name]
            _set_fonts(st.element, self.s.body_font)
            st.font.color.rgb = RGBColor(0, 0, 0)

    def _style_heading(self, st, hs: HeadingStyle) -> None:
        _set_fonts(st.element, hs.font)
        f = st.font
        f.size = Pt(hs.size)
        f.bold, f.italic, f.underline, f.all_caps = hs.bold, False, hs.underline, hs.caps
        f.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.alignment = _ALIGN[hs.align]
        pf.space_before, pf.space_after = Pt(hs.space_before), Pt(hs.space_after)
        pf.keep_with_next = True
        pf.line_spacing = 1.0

    # ---- cover page

    def add_cover(self, cover: CoverInfo) -> None:
        d = self.doc
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(54)
        p.paragraph_format.space_after = Pt(24)
        p.add_run().add_picture(str(self.profile.logo()), width=Inches(2.1))

        def caps_line(text: str, before: float, after: float = 6) -> None:
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(before)
            para.paragraph_format.space_after = Pt(after)
            lines = text.upper().split("\n")
            for i, line in enumerate(lines):
                r = para.add_run(line.strip())
                self._font(r, self.s.cover_font, 14, bold=True)
                if i < len(lines) - 1:
                    r.add_break()

        caps_line(self.profile.university, 12)
        caps_line(self.profile.department, 6)
        if cover.module:
            caps_line(cover.module, 30)

        t = d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_before = Pt(78)
        t.paragraph_format.space_after = Pt(0)
        t.paragraph_format.line_spacing = 1.5
        self._font(t.add_run(cover.title), self.s.title_font, self.s.title_size, bold=True)

        details = [
            ("Name", self.profile.name),
            ("Reg No", self.profile.reg_no),
            ("Intake", self.profile.intake),
            ("Submission Date", cover.submission_date),
        ]
        details = [(k, v) for k, v in details if v]
        for i, (label, value) in enumerate(details):
            para = d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Inches(0.4)
            pf.space_before = Pt(90 if i == 0 else 0)
            pf.space_after = Pt(10)
            pf.tab_stops.add_tab_stop(Inches(2.1))
            self._font(para.add_run(f"{label}\t"), self.s.cover_font, 12)
            r = para.add_run(value)
            self._font(r, self.s.cover_font, 12, bold=True)
            if i == len(details) - 1:
                r.add_break(WD_BREAK.PAGE)

    @staticmethod
    def _font(run, name: str, size: float, bold: bool = False) -> None:
        run.font.name = name
        _set_fonts(run._r, name)
        run.font.size = Pt(size)
        run.font.bold = bold

    # ---- body

    def add_markdown(self, markdown: str) -> None:
        ast = json.loads(
            pypandoc.convert_text(
                markdown,
                "json",
                format="markdown+tex_math_dollars+pipe_tables+grid_tables+fenced_code_blocks-implicit_figures",
            )
        )
        blocks = ast["blocks"]
        self._prepare_math(blocks)
        for b in blocks:
            self._block(self.doc, b, list_ctx=None)

    def _prepare_math(self, blocks) -> None:
        found: list[str] = []

        def walk(node):
            if isinstance(node, dict):
                if node.get("t") == "Math":
                    found.append(node["c"][1])
                    node["_mi"] = len(found) - 1
                else:
                    walk(node.get("c"))
            elif isinstance(node, list):
                for x in node:
                    walk(x)

        walk(blocks)
        if not found:
            return
        try:
            converted = latex_to_omml(found)
        except Exception as exc:  # pandoc failure: fall back to plain text equations
            self.warnings.append(f"Equation conversion failed ({exc}); equations inserted as text.")
            converted = [None] * len(found)
        for i, (tex, el) in enumerate(zip(found, converted)):
            if el is None:
                self.warnings.append(f"Could not convert equation to Word format: {tex}")
            self._math[i] = el

    # ---- blocks

    def _new_para(self, container, list_ctx, *, justify: bool = True):
        p = container.add_paragraph()
        if list_ctx is not None:
            num_id, ilvl, first = list_ctx["num"], list_ctx["lvl"], list_ctx["first"]
            if first:
                _apply_num(p, num_id, ilvl)
                list_ctx["first"] = False
            else:
                p.paragraph_format.left_indent = Inches(0.5 + 0.25 * ilvl)
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if list_ctx is None and container is self.doc and self._in_question():
            p.paragraph_format.left_indent = Inches(0.25)
        return p

    def _in_question(self) -> bool:
        # Assignment answers sit indented under their numbered question, as in the samples.
        return self.s.number_h1 and self._question_num is not None and not self._after_backmatter

    def _block(self, container, b, list_ctx) -> None:
        t, c = b["t"], b.get("c")
        if t in ("Para", "Plain"):
            self._para(container, c, list_ctx)
        elif t == "Header":
            self._header(container, c)
        elif t == "BulletList":
            self._list(container, c, None, list_ctx)
        elif t == "OrderedList":
            self._list(container, c[1], c[0], list_ctx)
        elif t == "Table":
            self._table(container, c)
        elif t == "Figure":
            self._figure(container, c)
        elif t == "CodeBlock":
            self._code(container, c[1])
        elif t == "BlockQuote":
            for sub in c:
                before = len(container.paragraphs)
                self._block(container, sub, list_ctx)
                for p in container.paragraphs[before:]:
                    p.paragraph_format.left_indent = Inches(0.5)
        elif t == "Div":
            for sub in c[1]:
                self._block(container, sub, list_ctx)
        elif t == "DefinitionList":
            self._deflist(container, c, list_ctx)
        elif t == "LineBlock":
            p = self._new_para(container, list_ctx, justify=False)
            for i, line in enumerate(c):
                self._inlines(p, line, {})
                if i < len(c) - 1:
                    p.add_run().add_break()
        elif t == "RawBlock":
            if _PAGEBREAK_RAW.match(c[1]):
                container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "HorizontalRule":
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _para(self, container, inlines, list_ctx) -> None:
        meaningful = [i for i in inlines if i["t"] not in ("Space", "SoftBreak", "LineBreak")]
        # A paragraph that is only an image is a figure.
        if len(meaningful) == 1 and meaningful[0]["t"] == "Image" and list_ctx is None:
            img = meaningful[0]["c"]
            self._figure_image(container, img, caption=img[1])
            return
        # A paragraph that is only display math is a centred equation.
        if len(meaningful) == 1 and meaningful[0]["t"] == "Math" and meaningful[0]["c"][0]["t"] == "DisplayMath":
            p = self._new_para(container, list_ctx, justify=False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._math_into(p, meaningful[0], display=True)
            return
        p = self._new_para(container, list_ctx)
        self._inlines(p, inlines, {})

    def _header(self, container, c) -> None:
        level, _attr, inlines = c
        level = min(level, 3)
        plain = "".join(x["c"] if x["t"] == "Str" else " " for x in inlines if x["t"] in ("Str", "Space")).strip()
        if level == 1 and self.s.number_h1 and _UNNUMBERED.match(plain):
            p = container.add_paragraph(style="Heading 2")  # "References" etc. are not questions
            self._after_backmatter = True
        elif level == 1 and self.s.number_h1:
            p = container.add_paragraph(style="Heading 1")
            if self._question_num is None:
                self._question_num = self.num.questions()
            _apply_num(p, self._question_num, 0)
            inlines = self._strip_question_prefix(inlines)
        else:
            p = container.add_paragraph(style=f"Heading {level}")
            if self._in_question():
                p.paragraph_format.left_indent = Inches(0.25)
        self._inlines(p, inlines, {})

    @staticmethod
    def _strip_question_prefix(inlines):
        # "# 1. Define a DAC" / "# Q1) Define a DAC" -> "Define a DAC" (numbering is automatic)
        for n in (1, 2):
            if len(inlines) > n and inlines[n]["t"] == "Space":
                words = [x["c"] if x["t"] == "Str" else " " for x in inlines[:n]]
                if all(x["t"] in ("Str", "Space") for x in inlines[:n]) and _QUESTION_PREFIX.fullmatch("".join(words) + " "):
                    return inlines[n + 1 :]
        return inlines

    def _list(self, container, items, attrs, list_ctx) -> None:
        level = 0 if list_ctx is None else list_ctx["lvl"] + 1
        if attrs is None:
            num_id = self.num.bullet()
        else:
            start, style, _delim = attrs
            num_id = self.num.ordered(style["t"], start)
        for item in items:
            ctx = {"num": num_id, "lvl": level, "first": True}
            for sub in item:
                self._block(container, sub, ctx)
        # breathing room after the whole list
        if list_ctx is None and container.paragraphs:
            container.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def _deflist(self, container, items, list_ctx) -> None:
        # "**Term** - definition" bullets, as used in the assignment samples.
        num_id = self.num.bullet()
        for term, defs in items:
            ctx = {"num": num_id, "lvl": 0 if list_ctx is None else list_ctx["lvl"] + 1, "first": True}
            p = self._new_para(container, ctx)
            self._inlines(p, term, {"bold": True})
            for d in defs:
                for j, blk in enumerate(d):
                    if j == 0 and blk["t"] in ("Para", "Plain"):
                        self._font_run(p.add_run(" - "), {"bold": True})
                        self._inlines(p, blk["c"], {})
                    else:
                        self._block(container, blk, ctx)

    # ---- inlines

    def _font_run(self, run, fmt) -> None:
        if fmt.get("bold"):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("underline"):
            run.underline = True
        if fmt.get("strike"):
            run.font.strike = True
        if fmt.get("sup"):
            run.font.superscript = True
        if fmt.get("sub"):
            run.font.subscript = True
        if fmt.get("smallcaps"):
            run.font.small_caps = True
        if fmt.get("code"):
            run.font.name = self.s.code_font
            _set_fonts(run._r, self.s.code_font)
            run.font.size = Pt(self.s.body_size - 1)

    def _inlines(self, p, inlines, fmt) -> None:
        buf: list[str] = []

        def flush():
            if buf:
                self._font_run(p.add_run("".join(buf)), fmt)
                buf.clear()

        for node in inlines:
            t, c = node["t"], node.get("c")
            if t == "Str":
                buf.append(c)
            elif t in ("Space", "SoftBreak"):
                buf.append(" ")
            elif t == "LineBreak":
                flush()
                p.add_run().add_break()
            elif t in ("Strong", "Emph", "Underline", "Strikeout", "Superscript", "Subscript", "SmallCaps"):
                flush()
                key = {
                    "Strong": "bold",
                    "Emph": "italic",
                    "Underline": "underline",
                    "Strikeout": "strike",
                    "Superscript": "sup",
                    "Subscript": "sub",
                    "SmallCaps": "smallcaps",
                }[t]
                self._inlines(p, c, {**fmt, key: True})
            elif t == "Code":
                flush()
                self._font_run(p.add_run(c[1]), {**fmt, "code": True})
            elif t == "Quoted":
                q = ("‘", "’") if c[0]["t"] == "SingleQuote" else ("“", "”")
                buf.append(q[0])
                flush()
                self._inlines(p, c[1], fmt)
                buf.append(q[1])
            elif t == "Math":
                flush()
                self._math_into(p, node, display=False)
            elif t in ("Link", "Span"):
                flush()
                self._inlines(p, c[1], fmt)
            elif t == "Cite":
                flush()
                self._inlines(p, c[1], fmt)
            elif t == "Image":
                flush()
                path = self._resolve(c[2][0])
                if path:
                    p.add_run().add_picture(str(path), width=self._img_width(path, c[0][2], inline=True))
            elif t == "RawInline":
                continue
            elif t == "Note":
                continue
        flush()

    def _math_into(self, p, node, display: bool) -> None:
        el = self._math.get(node.get("_mi", -1))
        if el is None:
            r = p.add_run(node["c"][1])
            r.italic = True
            return
        el = copy.deepcopy(el)
        if display:
            para = OxmlElement("m:oMathPara")
            para.append(el)
            p._p.append(para)
        else:
            p._p.append(el)

    # ---- figures / images

    def _resolve(self, src: str) -> Path | None:
        path = Path(src).expanduser()
        candidates = [path] if path.is_absolute() else []
        if not path.is_absolute():
            if self.assets_dir:
                candidates.append(self.assets_dir / path)
            candidates.append(Path.cwd() / path)
        for cand in candidates:
            if cand.is_file():
                return self._upright(cand)
        self.warnings.append(f"Image not found: {src}")
        return None

    def _upright(self, path: Path) -> Path:
        """Phone photos carry an EXIF rotation flag that Word does not always honour; bake it in."""
        try:
            with PILImage.open(path) as im:
                if im.getexif().get(0x0112, 1) in (1, None):
                    return path
                fixed = ImageOps.exif_transpose(im)
                out = Path(self._tmp.name) / f"{len(list(Path(self._tmp.name).iterdir()))}_{path.stem}.png"
                fixed.save(out)
                return out
        except Exception:
            return path

    def _img_width(self, path: Path, kvs, inline: bool = False):
        attrs = dict(kvs)
        width = attrs.get("width")
        if width:
            m = re.fullmatch(r"\s*([\d.]+)\s*(%|in|cm|mm|px)?\s*", width)
            if m:
                val, unit = float(m.group(1)), m.group(2) or "px"
                inches = {
                    "%": CONTENT_W * val / 100,
                    "in": val,
                    "cm": val / 2.54,
                    "mm": val / 25.4,
                    "px": val / 96,
                }[unit]
                return Inches(min(inches, CONTENT_W))
        with PILImage.open(path) as im:
            w_px, h_px = im.size
            dpi = im.info.get("dpi", (96, 96))[0] or 96
        w_in, h_in = w_px / dpi, h_px / dpi
        if inline:
            return Inches(min(w_in, 1.5))
        scale = min(1.0, CONTENT_W / w_in, MAX_IMG_H / h_in)
        # Small screenshots look lost on the page; bring tiny images up to at least half width.
        if w_in * scale < CONTENT_W * 0.5:
            scale = min(CONTENT_W * 0.5 / w_in, MAX_IMG_H / h_in)
        return Inches(w_in * scale)

    def _figure(self, container, c) -> None:
        _attr, caption, body = c
        cap_blocks = caption[1]
        cap_inlines = cap_blocks[0]["c"] if cap_blocks else []
        for blk in body:
            for node in blk.get("c", []) if blk["t"] in ("Para", "Plain") else []:
                if node["t"] == "Image":
                    self._figure_image(container, node["c"], caption=cap_inlines or node["c"][1])

    def _figure_image(self, container, img, caption) -> None:
        attr, _alt, (src, _title) = img
        path = self._resolve(src)
        if not path:
            p = container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[Missing image: {src}]")
            r.italic = True
            r.font.color.rgb = RGBColor(0xC0, 0, 0)
            return
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = self.s.figure_captions and bool(caption)
        p.add_run().add_picture(str(path), width=self._img_width(path, attr[2]))
        if self.s.figure_captions and caption:
            self._fig_n += 1
            cp = container.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(12)
            r = cp.add_run(f"Figure {self._fig_n}: ")
            r.bold = True
            r.font.size = Pt(self.s.body_size - 1)
            before = len(cp.runs)
            self._inlines(cp, caption, {"italic": True})
            for run in cp.runs[before:]:
                run.font.size = Pt(self.s.body_size - 1)

    # ---- tables

    def _table(self, container, c) -> None:
        _attr, caption, colspecs, head, bodies, foot = c
        rows = []
        for row in head[1]:
            rows.append((True, row[1]))
        for body in bodies:
            for row in body[2] + body[3]:
                rows.append((False, row[1]))
        for row in foot[1]:
            rows.append((False, row[1]))
        if not rows:
            return
        ncols = len(colspecs)

        cap_inlines = caption[1][0]["c"] if caption[1] else []
        if cap_inlines:
            self._tab_n += 1
            cp = container.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.keep_with_next = True
            r = cp.add_run(f"Table {self._tab_n}: ")
            r.bold = True
            self._inlines(cp, cap_inlines, {"italic": True})

        table = container.add_table(rows=len(rows), cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        widths = [spec[1].get("c") if spec[1]["t"] == "ColWidth" else None for spec in colspecs]
        if all(widths):
            total = sum(widths)
            widths = [CONTENT_W * 0.95 * w / total for w in widths]
        else:
            widths = [CONTENT_W * 0.95 / ncols] * ncols

        for r_i, (is_head, cells) in enumerate(rows):
            col = 0
            for cell in cells:
                _cattr, align, _rowspan, colspan, blocks = cell
                if col >= ncols:
                    break
                target = table.cell(r_i, col)
                if colspan > 1 and col + colspan - 1 < ncols:
                    target = target.merge(table.cell(r_i, col + colspan - 1))
                self._cell(target, blocks, is_head, align["t"])
                col += colspan
            if is_head:
                trpr = table.rows[r_i]._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                trpr.append(hdr)

        for row in table.rows:
            for i, cell in enumerate(row.cells[:ncols]):
                cell.width = Inches(widths[i])

        spacer = container.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _cell(self, cell, blocks, is_head: bool, align: str) -> None:
        first = cell.paragraphs[0]
        # Render into the cell, then move the first rendered paragraph's content into the
        # pre-existing empty paragraph so there is no blank line at the top of the cell.
        before = len(cell.paragraphs)
        for b in blocks:
            self._block(cell, b, None)
        new = cell.paragraphs[before:]
        if new:
            src = new[0]._p
            for child in list(src):
                first._p.append(child)
            src.getparent().remove(src)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.alignment = {
                "AlignCenter": WD_ALIGN_PARAGRAPH.CENTER,
                "AlignRight": WD_ALIGN_PARAGRAPH.RIGHT,
            }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
            if is_head and self.s.table_header_bold:
                for r in p.runs:
                    r.bold = True

    # ---- code

    def _code(self, container, text: str) -> None:
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(10), 1.0
        pf.left_indent = pf.right_indent = Inches(0.05)
        _para_box(p)
        lines = text.expandtabs(4).split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = self.s.code_font
            _set_fonts(r._r, self.s.code_font)
            r.font.size = Pt(self.s.code_size)
            if i < len(lines) - 1:
                r.add_break()

    # ---- output

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.doc.save(path)
        except PermissionError:
            # Usually the previous version is still open in Word.
            alt = path.with_name(f"{path.stem} (new){path.suffix}")
            self.doc.save(alt)
            self.warnings.append(f"{path.name} is open in another program; saved as {alt.name} instead.")
            return alt
        finally:
            self._tmp.cleanup()
        return path


def build_report(
    *,
    style: TemplateStyle,
    profile: Profile,
    cover: CoverInfo,
    markdown: str,
    out_path: Path,
    assets_dir: Path | None = None,
) -> BuildResult:
    b = ReportBuilder(style, profile, assets_dir)
    b.add_cover(cover)
    b.add_markdown(markdown)
    path = b.save(out_path)
    return BuildResult(path=path, warnings=b.warnings)


__all__ = ["CoverInfo", "BuildResult", "ReportBuilder", "build_report", "M"]
