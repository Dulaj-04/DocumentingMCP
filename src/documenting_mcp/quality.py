"""Pre-submission checks and a similarity check against source material."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import docx

# Words that are correct in EE reports but unknown to a general English dictionary.
EE_WORDS = set(
    """
    stm32 stm32cubeide nucleo hal i2c spi uart usart gpio adc dac pwm dma rtc mcu mcus soc fpga cpld
    vhdl verilog arduino esp32 microcontroller microcontrollers microprocessor op-amp opamp opamps mosfet
    mosfets bjt bjts igbt igbts jfet cmos ttl nmos pmos pcb pcbs lcd lcds led leds oled pmod pmodcls
    pscad matlab simulink multisim proteus ltspice spice orcad kicad labview psim plc plcs scada vfd vfds
    rms dc ac rpm pu kv kw kva mva mw ma ua hz khz mhz ghz db dbm ohm ohms kohm mohm nf pf uf mh uh
    datasheet datasheets firmware bootloader debounce debouncing setpoint setpoints thevenin norton
    sinusoidal non-linearity nonlinearity monotonicity ratiometric ratiometrically lsb msb lsbs msbs
    quantisation quantization oversampling sigma-delta resistive capacitive inductive impedance impedances
    reactance susceptance admittance phasor phasors transconductance transimpedance biasing optocoupler
    rectifier rectifiers thyristor thyristors triac triacs scr scrs inverter inverters chopper choppers
    buck boost flyback synchronous asynchronous rotor stator rotors stators slip-ring commutator
    multiplexer demultiplexer multiplexers flip-flop flip-flops latch latches decoder encoder encoders
    ladder vref vout vin iout pid feedforward overshoot undershoot bode nyquist laplace fourier z-transform
    transient transients steady-state eigenvalue eigenvalues wavelength bandwidth passband stopband
    kirchhoff kirchhoff's ohm's ampere's faraday's lenz's lorentz kdu
    """.split()
)
COMMON_TYPOS = {
    "disscussion": "discussion", "discusion": "discussion", "votage": "voltage", "voltge": "voltage",
    "recieve": "receive", "recieved": "received", "seperate": "separate", "occured": "occurred",
    "acheive": "achieve", "acheived": "achieved", "definately": "definitely", "refered": "referred",
    "transmiter": "transmitter", "freqency": "frequency", "frequncy": "frequency", "resistence": "resistance",
    "capacitence": "capacitance", "inductence": "inductance", "conclution": "conclusion", "experiement": "experiment",
    "apparatous": "apparatus", "proceedure": "procedure", "observaton": "observation", "calcualte": "calculate",
    "analysys": "analysis", "comparision": "comparison", "lenght": "length", "accross": "across",
    "untill": "until", "wich": "which", "teh": "the", "adress": "address", "succesfully": "successfully",
    "successfuly": "successfully", "implemention": "implementation", "intial": "initial", "intialise": "initialise",
}
_UNIT_NO_SPACE = re.compile(
    r"(?<![\w.])(\d+(?:\.\d+)?)(V|A|W|Hz|kHz|MHz|GHz|mA|µA|uA|kV|mV|kW|MW|VA|kVA|MVA|ms|µs|us|ns|Ω|kΩ|MΩ|mH|µH|µF|uF|nF|pF|dB|rpm|pu|°C)(?![\w])"
)
_PLACEHOLDER = re.compile(r"(\[(?:insert|add|todo|tbd|placeholder)[^\]]*\]|\bTODO\b|\bTBD\b|\bXXX\b|lorem ipsum|\[Missing image[^\]]*\])", re.I)
# Wording that makes a report read as AI-generated (lecturers notice these).
AI_PHRASES = {
    r"\bdelv(?:e|es|ed|ing)\b": "delve",
    r"\b(?:crucial|pivotal|paramount)\b": "crucial / pivotal / paramount",
    r"\bseamless(?:ly)?\b": "seamless",
    r"\brobust(?:ly|ness)?\b": "robust",
    r"\bleverag(?:e|es|ed|ing)\b": "leverage",
    r"\butili[sz](?:e|es|ed|ing|ation)\b": "utilise (say 'use')",
    r"\bharness(?:es|ed|ing)?\b": "harness",
    r"\bfoster(?:s|ed|ing)?\b": "foster",
    r"\b(?:landscape|realm|tapestry)\b": "landscape / realm / tapestry",
    r"\b(?:intricate|multifaceted|holistic)\b": "intricate / multifaceted / holistic",
    r"\bunderscor(?:e|es|ed|ing)\b": "underscore",
    r"\bshowcas(?:e|es|ed|ing)\b": "showcase",
    r"\bplays? an? (?:vital|crucial|key|pivotal|important) role\b": "'plays a vital role'",
    r"\bin today['’]?s (?:world|era|landscape)\b": "'in today's world'",
    r"\bit is (?:important|worth|crucial) (?:to note|noting)\b": "'it is worth noting that'",
    r"\bnot only\b[^.]{0,80}\bbut also\b": "'not only … but also'",
    r"\bin conclusion,": "'In conclusion,'",
    r"\b(?:a testament to|stands as)\b": "'a testament to'",
}
_FIRST_PERSON = re.compile(r"\b(?:I|[Ww]e|[Oo]ur|[Uu]s|[Mm]y)\b(?![-/0-9])")


def _american(word: str) -> list[str]:
    """American forms of a British spelling, so 'generalised', 'colour', 'centre' are accepted."""
    forms = []
    for gb, us in (("isation", "ization"), ("ising", "izing"), ("ised", "ized"), ("ises", "izes"), ("ise", "ize"),
                   ("yse", "yze"), ("ysed", "yzed"), ("ysing", "yzing"), ("our", "or"), ("ours", "ors"),
                   ("tre", "ter"), ("tres", "ters"), ("lled", "led"), ("lling", "ling"), ("ogue", "og"), ("ence", "ense")):
        if word.endswith(gb):
            forms.append(word[: -len(gb)] + us)
    if "our" in word:
        forms.append(word.replace("our", "or"))
    return forms


def _body_paragraphs(path: Path) -> list[tuple[str, str]]:
    """(style, text) for each paragraph after the cover page."""
    d = docx.Document(str(path))
    out, started = [], False
    for p in d.paragraphs:
        name = p.style.name if p.style is not None else ""
        if name.startswith("Heading"):
            started = True
        if started and p.text.strip():
            out.append((name, p.text.strip()))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    out.append(("Table", cell.text.strip()))
    return out


def _strip_md(md: str) -> str:
    md = re.sub(r"```.*?```", " ", md, flags=re.S)
    md = re.sub(r"\$\$.*?\$\$", " ", md, flags=re.S)
    md = re.sub(r"\$[^$\n]+\$", " ", md)
    md = re.sub(r"!\[[^\]]*\]\([^)]*\)(\{[^}]*\})?", " ", md)
    md = re.sub(r"`[^`]+`", " ", md)
    md = re.sub(r"\[@[^\]]+\]|@(?:fig|tbl|eq):[\w-]+", " ", md)
    return md


def check_document(
    docx_path: Path,
    markdown: str | None,
    build_warnings: list[str],
    *,
    word_limit: int | None = None,
    page_count: int | None = None,
    page_limit: int | None = None,
    expected_questions: int | None = None,
    template: str = "",
) -> dict:
    issues: list[dict] = []

    def add(sev: str, where: str, msg: str) -> None:
        issues.append({"severity": sev, "where": where, "message": msg})

    for w in build_warnings:
        add("error" if re.search(r"not found|missing|no matching|points to nothing|could not", w, re.I) else "warning", "build", w)

    paras = _body_paragraphs(docx_path)
    text = "\n".join(t for _, t in paras)
    words = re.findall(r"[A-Za-z][A-Za-z'’-]*", text)
    total_words = len(words)

    # Placeholders left behind
    for style, t in paras:
        for m in _PLACEHOLDER.finditer(t):
            add("error", t[:60], f"Placeholder left in the document: '{m.group(0)}'")

    # Spelling: known typos first, then a dictionary pass
    lower_counts = Counter(w.lower().strip("'’-") for w in words)
    for typo, fix in COMMON_TYPOS.items():
        if typo in lower_counts:
            add("error", typo, f"Spelling: '{typo}' should be '{fix}'")
    try:
        from spellchecker import SpellChecker

        sc = SpellChecker()
        sc.word_frequency.load_words(EE_WORDS)
        candidates = [w for w in lower_counts if len(w) >= 4 and w.isalpha() and w not in COMMON_TYPOS and w not in EE_WORDS]
        # Ignore proper nouns / acronyms: only check words that also appear in lower case in the text.
        lowercase_seen = {w for w in words if w.islower()}
        candidates = [w for w in candidates if w in lowercase_seen]
        for w in sorted(sc.unknown(candidates)):
            if any(f in sc for f in _american(w)):
                continue  # valid British spelling
            fix = sc.correction(w)
            if fix and fix != w:
                add("warning", w, f"Possible spelling mistake: '{w}' (did you mean '{fix}'?)")
    except Exception as exc:  # spell checker unavailable
        add("info", "spelling", f"Dictionary spell check skipped: {exc}")

    # Repeated words
    for m in re.finditer(r"\b(\w+)\s+\1\b", text, re.I):
        if m.group(1).lower() not in ("that", "had", "is"):
            add("warning", m.group(0), f"Repeated word: '{m.group(0)}'")

    # Units without a space (prose only; equations are typeset automatically)
    prose = _strip_md(markdown) if markdown else text
    for m in _UNIT_NO_SPACE.finditer(prose):
        add("warning", m.group(0), f"Put a space between number and unit: '{m.group(1)} {m.group(2)}'")

    # First person in a formal report
    fp = [t for _, t in paras if _FIRST_PERSON.search(t)]
    if fp:
        add("warning", fp[0][:60], f"First-person words (I/we/our) in {len(fp)} paragraph(s); formal reports use the passive voice.")

    # AI-style writing
    lowered = text.lower()
    for pattern, label in AI_PHRASES.items():
        hits = re.findall(pattern, lowered)
        if hits:
            add("warning", label, f"Sounds AI-generated: {label} used {len(hits)} time(s). Rewrite in plain words.")
    connectives = len(re.findall(r"\b(?:furthermore|moreover|additionally)\b", lowered))
    if connectives > 2:
        add("warning", "connectives", f"'Furthermore/Moreover/Additionally' used {connectives} times; this reads as AI text.")
    dashes = text.count("—")
    if total_words and dashes * 1000 / total_words > 4:
        add("warning", "em dashes", f"{dashes} em dashes (—); use commas, brackets or new sentences instead.")
    for _, t in paras:
        if re.search(r"\S\s+[·•|]\s+\S.*\s+[·•|]\s+\S", t) and ":" in t:
            add("warning", t[:60], "Decorative '·'-separated summary line; write it as a normal sentence or table.")
    d_ = docx.Document(str(docx_path))
    body = [p for p in d_.paragraphs if p.text.strip() and not (p.style is not None and p.style.name.startswith("Heading"))]
    labelled = [p for p in body if p.runs and p.runs[0].bold and re.match(r"^[^.]{2,40}:", p.text)]
    if len(body) >= 8 and len(labelled) / len(body) > 0.35:
        add("warning", "bold labels", f"{len(labelled)} of {len(body)} paragraphs start with a bold 'Label:'; vary the writing.")

    # Very long sentences
    for _, t in paras:
        for sent in re.split(r"(?<=[.!?])\s+", t):
            n = len(sent.split())
            if n > 45:
                add("info", sent[:60], f"Long sentence ({n} words); consider splitting it.")

    # Structure from the source
    if markdown:
        from .project import outline

        secs = outline(markdown)
        for k, s in enumerate(secs):
            nxt = secs[k + 1] if k + 1 < len(secs) else None
            topic_heading = template == "assignment" and s["level"] == 3 and nxt is not None  # "• Topic" + sub-headings
            generated = s["title"].strip().lower() in ("references", "bibliography")
            if s["words"] == 0 and not topic_heading and not generated and not (nxt and nxt["level"] > s["level"]):
                add("error", s["title"], "Section is empty.")
        labels = set(re.findall(r"\{#((?:fig|tbl|eq):[\w-]+)", markdown)) | set(re.findall(r"\\label\{((?:eq|fig|tbl):[^}]+)\}", markdown))
        for lab in sorted(labels):
            if f"@{lab}" not in markdown:
                add("info", lab, f"{lab} is never referred to in the text (use @{lab}).")
        if expected_questions is not None and template == "assignment":
            answered = sum(1 for s in secs if s["level"] == 1 and s["title"].strip().lower() not in ("references", "bibliography"))
            if answered < expected_questions:
                add("error", "questions", f"Only {answered} of {expected_questions} questions are answered.")
        if template == "lab_report":
            have = {s["title"].strip().lower() for s in secs if s["level"] == 1}
            for need in ("aim", "procedure", "result", "discussion", "conclusion"):
                if not any(h.startswith(need) for h in have):
                    add("warning", need, f"Lab report has no '{need.title()}' section.")

    if word_limit and total_words > word_limit:
        add("error", "length", f"{total_words} words, which is over the {word_limit}-word limit.")
    if page_limit and page_count and page_count > page_limit:
        add("error", "length", f"{page_count} pages, which is over the {page_limit}-page limit.")

    order = {"error": 0, "warning": 1, "info": 2}
    issues.sort(key=lambda i: order[i["severity"]])
    counts = Counter(i["severity"] for i in issues)
    return {
        "summary": {
            "words": total_words,
            "pages": page_count,
            "errors": counts.get("error", 0),
            "warnings": counts.get("warning", 0),
            "info": counts.get("info", 0),
            "ready_to_submit": counts.get("error", 0) == 0,
        },
        "issues": issues,
    }


# --------------------------------------------------------------------------- similarity


def _tokens(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def similarity_report(docx_path: Path, sources: dict[str, str], shingle: int = 6, threshold: float = 0.3) -> dict:
    """Flag paragraphs whose word 6-grams overlap heavily with any source text."""
    index: dict[tuple, tuple[str, int]] = {}
    source_tokens: dict[str, list[str]] = {}
    for name, text in sources.items():
        toks = _tokens(text)
        source_tokens[name] = toks
        for i in range(len(toks) - shingle + 1):
            index.setdefault(tuple(toks[i : i + shingle]), (name, i))

    flagged, total, matched = [], 0, 0
    for style, para in _body_paragraphs(docx_path):
        if style.startswith("Heading") or style == "Table":
            continue
        toks = _tokens(para)
        if len(toks) < 20:
            continue
        grams = [tuple(toks[i : i + shingle]) for i in range(len(toks) - shingle + 1)]
        hits = [index[g] for g in grams if g in index]
        total += len(grams)
        matched += len(hits)
        ratio = len(hits) / len(grams) if grams else 0
        if ratio >= threshold:
            name, pos = Counter(h[0] for h in hits).most_common(1)[0][0], None
            pos = next(h[1] for h in hits if h[0] == name)
            src = source_tokens[name]
            flagged.append({
                "paragraph": para[:300] + ("…" if len(para) > 300 else ""),
                "overlap": round(ratio * 100),
                "source": name,
                "source_excerpt": " ".join(src[max(0, pos - 5) : pos + 40]) + " …",
            })
    flagged.sort(key=lambda f: -f["overlap"])
    return {
        "overall_overlap_percent": round(100 * matched / total) if total else 0,
        "paragraphs_flagged": len(flagged),
        "threshold_percent": round(threshold * 100),
        "flagged": flagged,
        "advice": "Rewrite flagged paragraphs in your own words and cite the source with [@key]." if flagged else "No paragraph closely matches the sources.",
    }
