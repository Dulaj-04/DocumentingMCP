"""Render a report (cover details + Markdown body) to a .docx matching docs/style-spec.md.

The Markdown is parsed by pandoc into its JSON AST, then every block is written with python-docx
so that fonts, spacing, numbering, borders and equations are fully under our control.
"""

from __future__ import annotations

import copy
import json
import tempfile
import re
from dataclasses import dataclass, field
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import ImageOps

from . import fields
from .latex_prep import PreparedMath
from .latex_prep import prepare as prepare_math
from .omml import M, latex_to_omml
from .references import format_ieee
from .profile import Profile
from .styles import HeadingStyle, TemplateStyle

# Geometry measured from the master assignment: 0.5 in side margins, border ~32 pt from the page edge.
PAGE_W, PAGE_H = 8.5, 11.0
MARGIN_TB, MARGIN_SIDE = 0.7, 0.5
CONTENT_W = PAGE_W - 2 * MARGIN_SIDE
FIGURE_W = CONTENT_W * 0.92
# Optional "shaded" table style (navy header, banded rows); the default is plain Word "Table Grid".
TABLE_HEAD_FILL, TABLE_BAND_FILL, TABLE_BORDER = "2C3E7A", "EDF0F8", "BFBFBF"
MAX_IMG_H = 8.0
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}
_PAGEBREAK_RAW = re.compile(r"^\s*(<!--\s*page\s*break\s*-->|\\newpage|\\pagebreak|\\clearpage)\s*$", re.I)
_SECTION_NO = r"(?:\d+(?:\.\d+)*\.?\s+)?"  # optional "6. " / "6.1 " in front of a heading
_UNNUMBERED = re.compile(r"^" + _SECTION_NO + r"(references|bibliography|appendix\b.*|acknowledge?ments?)$", re.I)
_QUESTION_PREFIX = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?\d+\s*[.):]\s*", re.I)


@dataclass
class CoverInfo:
    title: str
    module: str = ""
    submission_date: str = ""
    department: str = ""  # overrides the profile's department for this document
    members: list = field(default_factory=list)  # group work: [{"name": ..., "reg_no": ...}]


@dataclass
class BuildOptions:
    references: dict[str, dict] = field(default_factory=dict)  # citation key -> IEEE entry
    figures_dir: Path | None = None  # where generated plots/diagrams/circuits are kept
    toc: bool = False
    list_of_figures: bool = False
    list_of_tables: bool = False
    page_numbers: str = "bottom-center"  # bottom-center (default) | bottom-right | none
    header_text: str = ""
    footer_text: str = ""
    approvals: dict[str, str] = field(default_factory=dict)  # figure key -> approved content digest
    border_width_pt: float = 0.5  # page border line width; 0 = no border (master: thin 0.5 pt)
    border_style: str = "single"  # single | double | thick-thin | dotted | dashed | none
    body_markdown: str = ""  # to tell whether a figure is referred to in the text


VISUAL_CLASSES = {
    "plot": "plot", "graph": "plot", "chart": "plot",
    "diagram": "diagram", "mermaid": "diagram", "flowchart": "diagram", "block": "diagram",
    "circuit": "circuit", "schematic": "circuit", "logic": "circuit",
}
_CROSSREF = re.compile(r"^(fig|tbl|tab|eq):", re.I)
_REFS_HEADING = re.compile(r"^" + _SECTION_NO + r"(references|bibliography|reference list)$", re.I)


@dataclass
class BuildResult:
    path: Path
    warnings: list[str] = field(default_factory=list)
    figures: list = field(default_factory=list)  # FigureRecord for every generated figure


# --------------------------------------------------------------------------- low-level XML helpers


def _set_fonts(rpr_owner, name: str) -> None:
    """Force a font on every script slot and drop theme-font overrides."""
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _shade(ppr_or_tcpr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr_or_tcpr.append(shd)


# Line widths Word offers for page borders, in points (stored as eighths of a point).
BORDER_WIDTHS_PT = (0.25, 0.5, 0.75, 1.0, 1.5, 2.25, 3.0, 4.5, 6.0)
BORDER_STYLES = {"single": "single", "double": "double", "thick-thin": "thickThinSmallGap", "dotted": "dotted", "dashed": "dashed"}


def _page_borders(sect_pr, width_pt: float = 0.5, style: str = "single") -> None:
    if width_pt <= 0 or style == "none":
        return
    width_pt = min(BORDER_WIDTHS_PT, key=lambda w: abs(w - width_pt))
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), BORDER_STYLES.get(style, "single"))
        el.set(qn("w:sz"), str(int(round(width_pt * 8))))
        el.set(qn("w:space"), "31")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    # pgBorders must come after pgMar/paperSrc in CT_SectPr; insert right after pgMar.
    pg_mar = sect_pr.find(qn("w:pgMar"))
    pg_mar.addnext(borders)


def _para_box(p) -> None:
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "BFBFBF")
        bdr.append(el)
    ppr.append(bdr)
    _shade(ppr, "F5F5F5")


# Child order required by the WordprocessingML schema (Word reports "unreadable content" otherwise).
_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
    "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
    "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
    "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
    "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
}


def _normalise_order(root) -> None:
    for local, seq in _ORDER.items():
        rank = {name: i for i, name in enumerate(seq.split())}
        for el in root.iter(qn(f"w:{local}")):
            children = list(el)
            children.sort(key=lambda c: rank.get(c.tag.split("}")[-1], len(rank)))
            for c in children:
                el.append(c)


def _set_col_widths(table, widths_in) -> None:
    """Set both the table grid and every cell; Word and LibreOffice read different ones."""
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        col.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row._tr.tc_lst, widths_in):
            tcpr = cell.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(int(w * 1440)))
            tcw.set(qn("w:type"), "dxa")
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tblw.set(qn("w:type"), "dxa")
    if tblpr.find(qn("w:tblLayout")) is None:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)


def _no_borders(table) -> None:
    """An invisible layout grid (explicit, so no table style can draw lines)."""
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _is_image_row(meaningful: list) -> bool:
    """A paragraph made only of two or more images is a row of figures placed side by side."""
    return len(meaningful) >= 2 and all(i["t"] == "Image" for i in meaningful)


def _plain_text(node) -> str:
    """Concatenated text of a pandoc AST fragment (for sizing table columns)."""
    if isinstance(node, list):
        return "".join(_plain_text(n) for n in node)
    if isinstance(node, dict):
        if node.get("t") == "Str":
            return node["c"]
        if node.get("t") in ("Space", "SoftBreak"):
            return " "
        if node.get("t") in ("Math", "Code"):
            return node["c"][1]
        return _plain_text(node.get("c"))
    return ""


# --------------------------------------------------------------------------- numbering


class Numbering:
    """Creates list definitions directly in numbering.xml so each list can restart at its own value."""

    BULLETS = ("•", "o", "▪")
    ORDERED_FMT = {
        "Decimal": ("decimal", "%{n}."),
        "DefaultStyle": ("decimal", "%{n}."),
        "Example": ("decimal", "%{n}."),
        "LowerRoman": ("lowerRoman", "%{n}."),
        "UpperRoman": ("upperRoman", "%{n}."),
        "LowerAlpha": ("lowerLetter", "%{n})"),
        "UpperAlpha": ("upperLetter", "%{n}."),
    }

    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        self._next_abs = 100
        self._next_num = 100

    def _lvl(self, ilvl: int, fmt: str, text: str, left: int, hanging: int, start: int = 1, font: str | None = None):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", str(start)), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            lvl.append(el)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        ppr.append(ind)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), font)
            rpr.append(rf)
            lvl.append(rpr)
        return lvl

    def _register(self, levels) -> int:
        abs_id, num_id = self._next_abs, self._next_num
        self._next_abs += 1
        self._next_num += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        abstract.append(mlt)
        for lvl in levels:
            abstract.append(lvl)
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            self.root.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num.append(ref)
        self.root.append(num)
        return num_id

    def bullet(self, base_left: int = 720) -> int:
        return self._register(
            [self._lvl(i, "bullet", self.BULLETS[i % 3], base_left + 360 * i, 360) for i in range(6)]
        )

    def ordered(self, style: str = "Decimal", start: int = 1, base_left: int = 720) -> int:
        fmt0, text0 = self.ORDERED_FMT.get(style, ("decimal", "%{n}."))
        nested = [("lowerLetter", "%{n})"), ("lowerRoman", "%{n}."), ("decimal", "%{n}.")]
        levels = [self._lvl(0, fmt0, text0.format(n=1), base_left, 360, start)]
        for i in range(1, 6):
            fmt, text = nested[(i - 1) % 3]
            levels.append(self._lvl(i, fmt, text.format(n=i + 1), base_left + 360 * i, 360))
        return self._register(levels)

    def questions(self) -> int:
        # Master: number at 0.125 in, question text at 0.375 in.
        return self._register([self._lvl(0, "decimal", "%1.", 540, 360)])


def _apply_num(p, num_id: int, ilvl: int) -> None:
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(lvl)
    numpr.append(nid)
    # numPr belongs near the start of pPr (after pStyle/keepNext etc.); python-docx tolerates insert at 0
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is not None:
        pstyle.addnext(numpr)
    else:
        ppr.insert(0, numpr)


# --------------------------------------------------------------------------- the builder


class ReportBuilder:
    def __init__(self, style: TemplateStyle, profile: Profile, assets_dir: Path | None = None, options: BuildOptions | None = None):
        self.opt = options or BuildOptions()
        self.s = style
        self.profile = profile
        self.assets_dir = assets_dir
        self.warnings: list[str] = []
        self.doc = Document()
        self.num = Numbering(self.doc)
        self._question_num: int | None = None
        self._after_backmatter = False
        self._h3_bullets: int | None = None
        self._math: dict[int, object] = {}
        self._math_meta: dict[int, PreparedMath] = {}
        self._eq_n = 0
        self._labels: dict[str, tuple[str, str]] = {}  # label -> (kind, number)
        self._captions: dict[str, list[str]] = {"Figure": [], "Table": []}
        self._cite_order: list[str] = []
        self._refs_header_block = None
        self._visual_n = 0
        self.figures: list = []
        self._front_placeholder = None
        self._tmp = tempfile.TemporaryDirectory()
        self._setup_page()
        self._setup_styles()
        self._setup_math()

    # ---- setup

    def _setup_page(self) -> None:
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(PAGE_W), Inches(PAGE_H)
        sec.top_margin = sec.bottom_margin = Inches(MARGIN_TB)
        sec.left_margin = sec.right_margin = Inches(MARGIN_SIDE)
        sec.header_distance = sec.footer_distance = Inches(0.5)
        _page_borders(sec._sectPr, self.opt.border_width_pt, self.opt.border_style)

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = self.s.body_font
        normal.font.size = Pt(self.s.body_size)
        _set_fonts(normal.element, self.s.body_font)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(self.s.para_after), 1.15

        for level, hs in ((1, self.s.h1), (2, self.s.h2), (3, self.s.h3)):
            st = styles[f"Heading {level}"]
            self._style_heading(st, hs)

        for name in ("Title", "Caption"):
            st = styles[name]
            _set_fonts(st.element, self.s.body_font)
            st.font.color.rgb = RGBColor(0, 0, 0)

    def _setup_math(self) -> None:
        """Word's document-level equation options: Cambria Math, display-style n-ary limits above/below."""
        settings = self.doc.settings.element
        old = settings.find(qn("m:mathPr"))
        math_pr = OxmlElement("m:mathPr")
        for tag, val in (
            ("m:mathFont", "Cambria Math"),
            ("m:brkBin", "before"),
            ("m:brkBinSub", "--"),
            ("m:smallFrac", "0"),
            ("m:dispDef", None),
            ("m:lMargin", "0"),
            ("m:rMargin", "0"),
            ("m:defJc", "centerGroup"),
            ("m:wrapIndent", "1440"),
            ("m:intLim", "subSup"),
            ("m:naryLim", "undOvr"),
        ):
            el = OxmlElement(tag)
            if val is not None:
                el.set(qn("m:val"), val)
            math_pr.append(el)
        if old is not None:
            old.addprevious(math_pr)
            old.getparent().remove(old)
        else:
            rsids = settings.find(qn("w:rsids"))
            (rsids.addnext if rsids is not None else settings.append)(math_pr)

    def _style_heading(self, st, hs: HeadingStyle) -> None:
        _set_fonts(st.element, hs.font)
        f = st.font
        f.size = Pt(hs.size)
        f.bold, f.italic, f.underline, f.all_caps = hs.bold, False, hs.underline, hs.caps
        f.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.alignment = _ALIGN[hs.align]
        pf.space_before, pf.space_after = Pt(hs.space_before), Pt(hs.space_after)
        pf.keep_with_next = True
        pf.line_spacing = 1.15

    # ---- cover page

    def add_cover(self, cover: CoverInfo) -> None:
        # Vertical rhythm copied from the master: crest top ≈ 1.2 in, university ≈ 3.2 in,
        # module ≈ 4.7 in, title ≈ 6.3 in, details from ≈ 7.9 in down the page.
        d = self.doc
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(36)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(self.profile.logo()), width=Inches(1.83))

        def caps_line(text: str, before: float, after: float = 11) -> None:
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(before)
            para.paragraph_format.space_after = Pt(after)
            lines = text.upper().split("\n")
            for i, line in enumerate(lines):
                r = para.add_run(line.strip())
                self._font(r, self.s.cover_heading_font, 14, bold=True)
                if i < len(lines) - 1:
                    r.add_break()

        caps_line(self.profile.university, 34)
        caps_line(cover.department or self.profile.department, 0)
        if cover.module:
            caps_line(cover.module, 28)

        t = d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_before = Pt(60 if cover.members else 84)
        t.paragraph_format.space_after = Pt(0)
        self._font(t.add_run(cover.title), self.s.title_font, self.s.title_size, bold=True)

        if cover.members:
            self._members_table(cover.members)
            details = [("Intake", self.profile.intake), ("Submission Date", cover.submission_date)]
        else:
            details = [
                ("Name", self.profile.name),
                ("Reg No", self.profile.reg_no),
                ("Intake", self.profile.intake),
                ("Submission Date", cover.submission_date),
            ]
        details = [(k, v) for k, v in details if v]
        for i, (label, value) in enumerate(details):
            para = d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Inches(0.25)
            pf.space_before = Pt((14 if cover.members else 86) if i == 0 else 0)
            pf.space_after = Pt(10)
            pf.tab_stops.add_tab_stop(Inches(1.75))
            self._font(para.add_run(f"{label}\t"), self.s.cover_font, 12)
            r = para.add_run(value)
            self._font(r, self.s.cover_font, 12, bold=True)
            if i == len(details) - 1 and not self._needs_body_section():
                r.add_break(WD_BREAK.PAGE)
        if self._needs_body_section():
            self._start_body_section()
        if self.opt.toc or self.opt.list_of_figures or self.opt.list_of_tables:
            self._front_placeholder = self.doc.add_paragraph()
            self._front_placeholder.add_run().add_break(WD_BREAK.PAGE)

    # ---- optional page furniture (off by default: the master has none)

    def _needs_body_section(self) -> bool:
        return self.opt.page_numbers != "none" or bool(self.opt.header_text or self.opt.footer_text)

    def _start_body_section(self) -> None:
        """New section after the cover so the cover carries no header/footer and numbering starts at 1."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sect_pr = sec._sectPr
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        anchor = sect_pr.find(qn("w:pgBorders"))
        if anchor is None:
            anchor = sect_pr.find(qn("w:pgMar"))
        anchor.addnext(pg)
        size = 10
        if self.opt.header_text:
            sec.header.is_linked_to_previous = False
            hp = sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp.add_run(self.opt.header_text)
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        if self.opt.footer_text or self.opt.page_numbers != "none":
            sec.footer.is_linked_to_previous = False
            fp = sec.footer.paragraphs[0]
            tabs = fp.paragraph_format.tab_stops
            for default in (3.25, 6.5):  # the Footer style's own tab stops would win otherwise
                tabs.add_tab_stop(Inches(default), WD_TAB_ALIGNMENT.CLEAR)
            tabs.add_tab_stop(Inches(CONTENT_W / 2), WD_TAB_ALIGNMENT.CENTER)
            tabs.add_tab_stop(Inches(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
            if self.opt.footer_text:
                fp.add_run(self.opt.footer_text).font.size = Pt(size)
            fp.add_run("\t")
            if self.opt.page_numbers == "bottom-center":
                fields.add_field(fp, "PAGE", "1", size_pt=size)
            fp.add_run("\t")
            if self.opt.page_numbers == "bottom-right":
                fields.add_field(fp, "PAGE", "1", size_pt=size)

    def _members_table(self, members: list) -> None:
        """Group cover: a plain No. / Name / Reg No table in place of the single-student lines."""
        d = self.doc
        lead = d.add_paragraph()
        lead.paragraph_format.space_before = Pt(40)
        lead.paragraph_format.space_after = Pt(6)
        lead.paragraph_format.left_indent = Inches(0.25)
        self._font(lead.add_run("Group Members"), self.s.cover_font, 12, bold=True)
        table = d.add_table(rows=len(members) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "000000")):
                el.set(qn(k), v)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        rows = [("No.", "Name", "Reg No")] + [
            (str(i), str(m.get("name", "")), str(m.get("reg_no", ""))) for i, m in enumerate(members, 1)
        ]
        for r_i, values in enumerate(rows):
            for c_i, val in enumerate(values):
                para = table.cell(r_i, c_i).paragraphs[0]
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                self._font(para.add_run(val), self.s.cover_font, 11, bold=(r_i == 0))
        _set_col_widths(table, (0.6, 3.4, 2.4))

    @staticmethod
    def _font(run, name: str, size: float, bold: bool = False) -> None:
        run.font.name = name
        _set_fonts(run._r, name)
        run.font.size = Pt(size)
        run.font.bold = bold

    # ---- body

    def add_markdown(self, markdown: str) -> None:
        ast = json.loads(
            pypandoc.convert_text(
                markdown,
                "json",
                format="markdown+tex_math_dollars+pipe_tables+grid_tables+fenced_code_blocks-implicit_figures",
            )
        )
        blocks = ast["blocks"]
        self._prepass(blocks)
        self._prepare_math(blocks)
        blocks = self._place_reference_list(blocks)
        for b in blocks:
            self._block(self.doc, b, list_ctx=None)
            if b is self._refs_header_block:
                self._reference_list()
        if self._refs_header_block is None and self._reference_keys():
            heading = {"t": "Header", "c": [1, ["", [], []], [{"t": "Str", "c": "References"}]]}
            self._block(self.doc, heading, list_ctx=None)
            self._reference_list()
        self._front_matter()

    # ---- numbering pre-pass (figures, tables, citations) so cross-references can point forward

    def _prepass(self, blocks) -> None:
        fig_n = tab_n = 0

        def figure(node_holder: dict, has_caption: bool, ident: str) -> None:
            nonlocal fig_n
            if (self.s.figure_captions and has_caption) or ident.lower().startswith("fig:"):
                fig_n += 1
                node_holder["_fig_no"] = fig_n
                if ident:
                    self._labels[ident] = ("Figure", str(fig_n))

        def walk(node):
            nonlocal tab_n
            if isinstance(node, list):
                for x in node:
                    walk(x)
                return
            if not isinstance(node, dict):
                return
            t = node.get("t")
            if t == "Para":
                meaningful = [i for i in node["c"] if i["t"] not in ("Space", "SoftBreak", "LineBreak")]
                if len(meaningful) == 1 and meaningful[0]["t"] == "Image":
                    img = meaningful[0]
                    figure(img, bool(img["c"][1]), img["c"][0][0])
                elif _is_image_row(meaningful):
                    first = meaningful[0]["c"][0]
                    figure(node, bool(dict(first[2]).get("caption")), first[0])
            elif t == "CodeBlock":
                attr = node["c"][0]
                if any(c.lower() in VISUAL_CLASSES for c in attr[1]):
                    kv = dict(attr[2])
                    figure(node, bool(kv.get("caption")), attr[0])
                    if kv.get("caption"):  # captions may use *emphasis* and $maths$
                        parsed = json.loads(pypandoc.convert_text(kv["caption"], "json", format="markdown+tex_math_dollars"))
                        blocks_ = parsed.get("blocks") or []
                        node["_cap"] = blocks_[0]["c"] if blocks_ else []
            elif t == "Table":
                attr, caption = node["c"][0], node["c"][1]
                if caption[1] or attr[0].lower().startswith(("tbl:", "tab:")):
                    tab_n += 1
                    node["_tab_no"] = tab_n
                    if attr[0]:
                        self._labels[attr[0]] = ("Table", str(tab_n))
            elif t == "Cite":
                for cit in node["c"][0]:
                    key = cit["citationId"]
                    if not _CROSSREF.match(key) and key not in self._cite_order:
                        self._cite_order.append(key)
            walk(node.get("c"))

        walk(blocks)

    def _reference_keys(self) -> list[str]:
        if self._cite_order:
            return self._cite_order
        return list(self.opt.references)  # nothing cited: list everything supplied

    def _place_reference_list(self, blocks):
        """If the body has a References heading and we generate the list, drop any hand-written list under it."""
        if not self._reference_keys():
            return blocks
        out, skipping, level = [], False, None
        for b in blocks:
            if b["t"] == "Header":
                text = _plain_text(b["c"][2]).strip()
                if skipping and b["c"][0] <= level:
                    skipping = False
                if self._refs_header_block is None and _REFS_HEADING.match(text):
                    self._refs_header_block, skipping, level = b, True, b["c"][0]
                    out.append(b)
                    continue
            if skipping:
                if b["t"] in ("OrderedList", "BulletList", "Para"):
                    self.warnings.append("A hand-written reference list was replaced by the generated IEEE list.")
                continue
            out.append(b)
        return out

    def _reference_list(self) -> None:
        keys = self._reference_keys()
        for n, key in enumerate(keys, 1):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(0.45), Inches(-0.45)
            pf.tab_stops.add_tab_stop(Inches(0.45))
            pf.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            bid = fields.bookmark_start(p, f"_Ref_cite{n}")
            p.add_run(f"[{n}]")
            fields.bookmark_end(p, bid)
            p.add_run("\t")
            entry = self.opt.references.get(key)
            if entry is None:
                self.warnings.append(f"Citation [@{key}] has no matching reference entry.")
                r = p.add_run(f"(missing reference: {key})")
                r.font.color.rgb = RGBColor(0xC0, 0, 0)
                continue
            for text, italic in format_ieee(entry):
                r = p.add_run(text)
                r.italic = italic

    # ---- table of contents / lists of figures and tables (opt-in)

    def _front_matter(self) -> None:
        ph = self._front_placeholder
        if ph is None:
            return
        parts = []
        if self.opt.toc:
            entries = []
            for para in self.doc.paragraphs:
                name = para.style.name if para.style is not None else ""
                if name.startswith("Heading ") and para.text.strip():
                    entries.append((int(name.split()[-1]), para.text.strip()))
            parts.append(("Table of Contents", 'TOC \\o "1-3" \\h \\z \\u', entries))
        if self.opt.list_of_figures:
            parts.append(("List of Figures", 'TOC \\h \\z \\c "Figure"', [(1, c) for c in self._captions["Figure"]]))
        if self.opt.list_of_tables:
            parts.append(("List of Tables", 'TOC \\h \\z \\c "Table"', [(1, c) for c in self._captions["Table"]]))
        for title, instr, entries in parts:
            tp = OxmlElement("w:p")
            ph._p.addprevious(tp)
            from docx.text.paragraph import Paragraph

            t = Paragraph(tp, ph._parent)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t.paragraph_format.space_before = Pt(6)
            t.paragraph_format.space_after = Pt(14)
            r = t.add_run(title)
            r.bold, r.font.size = True, Pt(14)
            body = OxmlElement("w:p")
            ph._p.addprevious(body)
            bp = Paragraph(body, ph._parent)
            cached = "\n".join(("    " * (lvl - 1)) + text for lvl, text in entries) or "Update this field in Word (right-click → Update Field)."
            fields.add_field(bp, instr, cached.split("\n")[0])
            # remaining cached lines as line breaks inside the result keep the field in one paragraph
            for line in cached.split("\n")[1:]:
                rr = OxmlElement("w:r")
                rr.append(OxmlElement("w:br"))
                tt = OxmlElement("w:t")
                tt.text = line
                tt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                rr.append(tt)
                end = [e for e in bp._p.iter(qn("w:fldChar")) if e.get(qn("w:fldCharType")) == "end"][-1].getparent()
                end.addprevious(rr)
            bp.paragraph_format.space_after = Pt(18)
        fields.set_update_fields_on_open(self.doc)

    def _prepare_math(self, blocks) -> None:
        found: list[str] = []

        def walk(node):
            if isinstance(node, dict):
                if node.get("t") == "Math":
                    prepared = prepare_math(node["c"][1], scientific=self.s.scientific_math)
                    if node["c"][0]["t"] == "DisplayMath" and not prepared.notag:
                        if prepared.tag:
                            prepared.number = prepared.tag
                        elif self.s.number_equations or prepared.label:
                            self._eq_n += 1
                            prepared.number = str(self._eq_n)
                        if prepared.label and prepared.number:
                            self._labels[prepared.label] = ("Equation", prepared.number)
                    self._math_meta[len(found)] = prepared
                    found.append(prepared.latex)
                    node["_mi"] = len(found) - 1
                else:
                    walk(node.get("c"))
                    walk(node.get("_cap"))
            elif isinstance(node, list):
                for x in node:
                    walk(x)

        walk(blocks)
        if not found:
            return
        try:
            converted = latex_to_omml(found)
        except Exception as exc:  # pandoc failure: fall back to plain text equations
            self.warnings.append(f"Equation conversion failed ({exc}); equations inserted as text.")
            converted = [None] * len(found)
        for i, (tex, el) in enumerate(zip(found, converted)):
            if el is None:
                self.warnings.append(f"Could not convert equation to Word format: {tex}")
            self._math[i] = el

    # ---- blocks

    def _new_para(self, container, list_ctx, *, justify: bool = True):
        p = container.add_paragraph()
        if list_ctx is not None:
            num_id, ilvl, first = list_ctx["num"], list_ctx["lvl"], list_ctx["first"]
            if first:
                _apply_num(p, num_id, ilvl)
                list_ctx["first"] = False
            else:
                p.paragraph_format.left_indent = Inches(list_ctx.get("left", self.s.list_text_indent) + 0.25 * ilvl)
            p.paragraph_format.space_after = Pt(self.s.list_item_after)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify and self.s.justify else WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify and self.s.justify else WD_ALIGN_PARAGRAPH.LEFT
        if list_ctx is None and container is self.doc and self._in_question():
            p.paragraph_format.left_indent = Inches(self.s.body_indent)
        return p

    def _in_question(self) -> bool:
        # Assignment answers sit indented under their numbered question, as in the samples.
        return self.s.number_h1 and self._question_num is not None and not self._after_backmatter

    def _block(self, container, b, list_ctx) -> None:
        t, c = b["t"], b.get("c")
        if t in ("Para", "Plain"):
            self._para(container, c, list_ctx, number=b.get("_fig_no"))
        elif t == "Header":
            self._header(container, c)
        elif t == "BulletList":
            self._list(container, c, None, list_ctx)
        elif t == "OrderedList":
            self._list(container, c[1], c[0], list_ctx)
        elif t == "Table":
            self._table(container, c, number=b.get("_tab_no"))
        elif t == "Figure":
            self._figure(container, c)
        elif t == "CodeBlock":
            if any(cl.lower() in VISUAL_CLASSES for cl in c[0][1]):
                self._visual(container, b)
            else:
                self._code(container, c[1])
        elif t == "BlockQuote":
            for sub in c:
                before = len(container.paragraphs)
                self._block(container, sub, list_ctx)
                for p in container.paragraphs[before:]:
                    p.paragraph_format.left_indent = Inches(0.5)
        elif t == "Div":
            for sub in c[1]:
                self._block(container, sub, list_ctx)
        elif t == "DefinitionList":
            self._deflist(container, c, list_ctx)
        elif t == "LineBlock":
            p = self._new_para(container, list_ctx, justify=False)
            for i, line in enumerate(c):
                self._inlines(p, line, {})
                if i < len(c) - 1:
                    p.add_run().add_break()
        elif t == "RawBlock":
            if _PAGEBREAK_RAW.match(c[1]):
                container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "HorizontalRule":
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _para(self, container, inlines, list_ctx, number=None) -> None:
        meaningful = [i for i in inlines if i["t"] not in ("Space", "SoftBreak", "LineBreak")]
        # A paragraph that is only an image is a figure.
        if len(meaningful) == 1 and meaningful[0]["t"] == "Image" and list_ctx is None:
            img = meaningful[0]
            self._figure_image(container, img["c"], caption=img["c"][1], number=img.get("_fig_no"))
            return
        # Two or more images and nothing else: a row of figures side by side.
        if _is_image_row(meaningful) and list_ctx is None:
            self._image_row(container, meaningful, number=number)
            return
        # Display math is always set on its own centred line, even when written mid-paragraph.
        if any(i["t"] == "Math" and i["c"][0]["t"] == "DisplayMath" for i in meaningful):
            segment: list = []
            for node in inlines + [None]:
                if node is None or (node["t"] == "Math" and node["c"][0]["t"] == "DisplayMath"):
                    if any(x["t"] not in ("Space", "SoftBreak", "LineBreak") for x in segment):
                        self._inlines(self._new_para(container, list_ctx), segment, {})
                    segment = []
                    if node is not None:
                        self._display_equation(container, node, list_ctx)
                else:
                    segment.append(node)
            return
        p = self._new_para(container, list_ctx)
        self._inlines(p, inlines, {})

    def _display_equation(self, container, node, list_ctx) -> None:
        meta = self._math_meta.get(node.get("_mi", -1))
        number = meta.number if meta else None
        if number is None or container is not self.doc:
            p = self._new_para(container, list_ctx, justify=False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # centred within the answer column, as in the master
            p.paragraph_format.space_before = Pt(4)
            self._math_into(p, node, display=True)
            return
        # Numbered equation: borderless 3-column row, equation centred on the page, "(n)" flush right.
        table = container.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (CONTENT_W * 0.12, CONTENT_W * 0.76, CONTENT_W * 0.12)
        _set_col_widths(table, widths)
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        mid = table.cell(0, 1).paragraphs[0]
        mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._math_into(mid, node, display=True)
        right = table.cell(0, 2).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        bid = fields.bookmark_start(right, fields.bookmark_name(meta.label)) if meta.label else None
        right.add_run("(")
        if number.isdigit():
            fields.add_field(right, "SEQ Equation \\* ARABIC", number)
        else:
            right.add_run(number)
        right.add_run(")")
        if bid:
            fields.bookmark_end(right, bid)
        spacer = container.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 0.5

    def _header(self, container, c) -> None:
        level, _attr, inlines = c
        level = min(level, 3)
        plain = "".join(x["c"] if x["t"] == "Str" else " " for x in inlines if x["t"] in ("Str", "Space")).strip()
        if level == 1 and self.s.number_h1 and _UNNUMBERED.match(plain):
            p = container.add_paragraph(style="Heading 2")  # "References" etc. are not questions
            self._after_backmatter = True
        elif level == 1 and self.s.number_h1:
            p = container.add_paragraph(style="Heading 1")
            if self._question_num is None:
                self._question_num = self.num.questions()
            _apply_num(p, self._question_num, 0)
            inlines = self._strip_question_prefix(inlines)
        else:
            p = container.add_paragraph(style=f"Heading {level}")
            if level == 3 and self.s.h3_bulleted:
                # "• Binary Weighted Resistor DAC": bullet at 0.25 in, text at 0.5 in.
                if self._h3_bullets is None:
                    self._h3_bullets = self.num.bullet(base_left=720)
                _apply_num(p, self._h3_bullets, 0)
        self._inlines(p, inlines, {})

    @staticmethod
    def _strip_question_prefix(inlines):
        # "# 1. Define a DAC" / "# Q1) Define a DAC" -> "Define a DAC" (numbering is automatic)
        for n in (1, 2):
            if len(inlines) > n and inlines[n]["t"] == "Space":
                words = [x["c"] if x["t"] == "Str" else " " for x in inlines[:n]]
                if all(x["t"] in ("Str", "Space") for x in inlines[:n]) and _QUESTION_PREFIX.fullmatch("".join(words) + " "):
                    return inlines[n + 1 :]
        return inlines

    def _list(self, container, items, attrs, list_ctx) -> None:
        level = 0 if list_ctx is None else list_ctx["lvl"] + 1
        # Lists inside table cells start near the cell edge rather than at the page list indent.
        left = 0.25 if container is not self.doc else self.s.list_text_indent
        base = int(left * 1440)
        if attrs is None:
            num_id = self.num.bullet(base_left=base)
        else:
            start, style, _delim = attrs
            num_id = self.num.ordered(style["t"], start, base_left=base)
        for item in items:
            ctx = {"num": num_id, "lvl": level, "first": True, "left": left}
            for sub in item:
                self._block(container, sub, ctx)
        # breathing room after the whole list
        if list_ctx is None and container.paragraphs:
            container.paragraphs[-1].paragraph_format.space_after = Pt(self.s.para_after)

    def _deflist(self, container, items, list_ctx) -> None:
        # "**Term** - definition" bullets, as used in the assignment samples.
        num_id = self.num.bullet()
        for term, defs in items:
            ctx = {"num": num_id, "lvl": 0 if list_ctx is None else list_ctx["lvl"] + 1, "first": True}
            p = self._new_para(container, ctx)
            self._inlines(p, term, {"bold": True})
            for d in defs:
                for j, blk in enumerate(d):
                    if j == 0 and blk["t"] in ("Para", "Plain"):
                        self._font_run(p.add_run(" - "), {"bold": True})
                        self._inlines(p, blk["c"], {})
                    else:
                        self._block(container, blk, ctx)

    # ---- inlines

    def _font_run(self, run, fmt) -> None:
        if fmt.get("bold"):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("underline"):
            run.underline = True
        if fmt.get("strike"):
            run.font.strike = True
        if fmt.get("sup"):
            run.font.superscript = True
        if fmt.get("sub"):
            run.font.subscript = True
        if fmt.get("smallcaps"):
            run.font.small_caps = True
        if fmt.get("code"):
            run.font.name = self.s.code_font
            _set_fonts(run._r, self.s.code_font)
            run.font.size = Pt(self.s.body_size - 1)

    def _inlines(self, p, inlines, fmt) -> None:
        buf: list[str] = []

        def flush():
            if buf:
                self._font_run(p.add_run("".join(buf)), fmt)
                buf.clear()

        for node in inlines:
            t, c = node["t"], node.get("c")
            if t == "Str":
                buf.append(c)
            elif t in ("Space", "SoftBreak"):
                buf.append(" ")
            elif t == "LineBreak":
                flush()
                p.add_run().add_break()
            elif t in ("Strong", "Emph", "Underline", "Strikeout", "Superscript", "Subscript", "SmallCaps"):
                flush()
                key = {
                    "Strong": "bold",
                    "Emph": "italic",
                    "Underline": "underline",
                    "Strikeout": "strike",
                    "Superscript": "sup",
                    "Subscript": "sub",
                    "SmallCaps": "smallcaps",
                }[t]
                self._inlines(p, c, {**fmt, key: True})
            elif t == "Code":
                flush()
                self._font_run(p.add_run(c[1]), {**fmt, "code": True})
            elif t == "Quoted":
                q = ("‘", "’") if c[0]["t"] == "SingleQuote" else ("“", "”")
                buf.append(q[0])
                flush()
                self._inlines(p, c[1], fmt)
                buf.append(q[1])
            elif t == "Math":
                flush()
                self._math_into(p, node, display=False)
            elif t in ("Link", "Span"):
                flush()
                self._inlines(p, c[1], fmt)
            elif t == "Cite":
                flush()
                self._cite(p, c[0], fmt)
            elif t == "Image":
                flush()
                path = self._resolve(c[2][0])
                if path:
                    p.add_run().add_picture(str(path), width=self._img_width(path, c[0][2], inline=True))
            elif t == "RawInline":
                continue
            elif t == "Note":
                continue
        flush()

    def _cite(self, p, citations, fmt) -> None:
        """[@key] -> IEEE [n]; @fig:x / @tbl:x / @eq:x -> live cross-reference to the caption or number."""
        nums: list[tuple[int, str]] = []

        def flush_nums():
            if not nums:
                return
            groups, run = [], [nums[0]]
            for n in nums[1:]:
                if n[0] == run[-1][0] + 1 and not n[1] and not run[-1][1]:
                    run.append(n)
                else:
                    groups.append(run)
                    run = [n]
            groups.append(run)
            for gi, g in enumerate(groups):
                if gi:
                    p.add_run(", ")
                if len(g) >= 3:
                    fields.internal_link(p, f"_Ref_cite{g[0][0]}", f"[{g[0][0]}]")
                    p.add_run("–")
                    fields.internal_link(p, f"_Ref_cite{g[-1][0]}", f"[{g[-1][0]}]")
                else:
                    for k, (n, suffix) in enumerate(g):
                        if k:
                            p.add_run(", ")
                        fields.internal_link(p, f"_Ref_cite{n}", f"[{n}{suffix}]")
            nums.clear()

        prev_crossref = False
        for i, cit in enumerate(citations):
            key = cit["citationId"]
            if _CROSSREF.match(key):
                flush_nums()
                if i:
                    p.add_run(", ")
                prev_crossref = True
                kind, number = self._labels.get(key, (None, None))
                if kind is None:
                    self.warnings.append(f"Cross-reference @{key} points to nothing (add {{#{key}}} to a figure, table or \\label{{{key}}} to an equation).")
                    r = p.add_run(f"??{key}")
                    r.font.color.rgb = RGBColor(0xC0, 0, 0)
                    continue
                bm = fields.bookmark_name(key)
                if kind == "Equation":
                    p.add_run("Eq. ")
                    fields.add_field(p, f"REF {bm} \\h", f"({number})")
                else:
                    fields.add_field(p, f"REF {bm} \\h", f"{kind} {number}")
                continue
            n = self._cite_order.index(key) + 1 if key in self._cite_order else 0
            suffix = _plain_text(cit.get("citationSuffix", [])).strip()
            if suffix and not suffix.startswith(","):
                suffix = ", " + suffix
            if prev_crossref:
                p.add_run(" ")
                prev_crossref = False
            if n == 0:
                p.add_run("[?]")
                continue
            nums.append((n, suffix))
        flush_nums()

    def _math_into(self, p, node, display: bool) -> None:
        el = self._math.get(node.get("_mi", -1))
        if el is None:
            r = p.add_run(node["c"][1])
            r.italic = True
            return
        el = copy.deepcopy(el)
        if display:
            para = OxmlElement("m:oMathPara")
            para.append(el)
            p._p.append(para)
        else:
            p._p.append(el)

    # ---- figures / images

    def _resolve(self, src: str) -> Path | None:
        path = Path(src).expanduser()
        candidates = [path] if path.is_absolute() else []
        if not path.is_absolute():
            if self.assets_dir:
                candidates.append(self.assets_dir / path)
            candidates.append(Path.cwd() / path)
        for cand in candidates:
            if cand.is_file():
                return self._upright(cand)
        from . import images as image_library

        found = image_library.find(path.name if path.is_absolute() else str(path))
        if found:
            return self._upright(found)
        self.warnings.append(f"Image not found: {src}")
        return None

    def _upright(self, path: Path) -> Path:
        """Phone photos carry an EXIF rotation flag that Word does not always honour; bake it in."""
        try:
            with PILImage.open(path) as im:
                if im.getexif().get(0x0112, 1) in (1, None):
                    return path
                fixed = ImageOps.exif_transpose(im)
                out = Path(self._tmp.name) / f"{len(list(Path(self._tmp.name).iterdir()))}_{path.stem}.png"
                fixed.save(out)
                return out
        except Exception:
            return path

    def _img_width(self, path: Path, kvs, inline: bool = False):
        attrs = dict(kvs)
        width = attrs.get("width")
        if width:
            m = re.fullmatch(r"\s*([\d.]+)\s*(%|in|cm|mm|px)?\s*", width)
            if m:
                val, unit = float(m.group(1)), m.group(2) or "px"
                inches = {
                    "%": FIGURE_W * val / 100,
                    "in": val,
                    "cm": val / 2.54,
                    "mm": val / 25.4,
                    "px": val / 96,
                }[unit]
                return Inches(min(inches, CONTENT_W))
        with PILImage.open(path) as im:
            w_px, h_px = im.size
            dpi = im.info.get("dpi", (96, 96))[0] or 96
        w_in, h_in = w_px / dpi, h_px / dpi
        if inline:
            return Inches(min(w_in, 1.5))
        scale = min(1.0, FIGURE_W / w_in, MAX_IMG_H / h_in)
        # Small screenshots look lost on the page; bring tiny images up to at least half width.
        if w_in * scale < FIGURE_W * 0.5:
            scale = min(FIGURE_W * 0.5 / w_in, MAX_IMG_H / h_in)
        return Inches(w_in * scale)

    def _figure(self, container, c) -> None:
        _attr, caption, body = c
        cap_blocks = caption[1]
        cap_inlines = cap_blocks[0]["c"] if cap_blocks else []
        for blk in body:
            for node in blk.get("c", []) if blk["t"] in ("Para", "Plain") else []:
                if node["t"] == "Image":
                    self._figure_image(container, node["c"], caption=cap_inlines or node["c"][1])

    def _figure_image(self, container, img, caption, number=None) -> None:
        attr, _alt, (src, _title) = img
        path = self._resolve(src)
        if not path:
            p = container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[Missing image: {src}]")
            r.italic = True
            r.font.color.rgb = RGBColor(0xC0, 0, 0)
            return
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = number is not None
        p.add_run().add_picture(str(path), width=self._img_width(path, attr[2]))
        if number is not None:
            self._caption(container, "Figure", number, attr[0], caption)

    ROW_MAX_COLS = 3
    ROW_MAX_H = 3.4  # inches: keeps a row from filling a whole page

    def _image_row(self, container, images, number) -> None:
        """Images side by side in an invisible grid: one common height, centred, with (a), (b)... sub-captions.

        Each image's alt text is its sub-caption. The first image may carry caption="..." and an id for the
        figure as a whole. Widths given with {width=...} are honoured; otherwise every image in a line gets the
        same height so the row looks tidy."""
        n = len(images)
        cols = min(n, self.ROW_MAX_COLS) if n <= 4 and n != 4 else (2 if n == 4 else self.ROW_MAX_COLS)
        grid = [images[i : i + cols] for i in range(0, n, cols)]
        gap = 0.15
        col_w = (CONTENT_W - gap * (cols - 1)) / cols
        usable = col_w - 0.22  # Word keeps 0.08" of padding inside each cell; a wider picture gets clipped
        table = container.add_table(rows=len(grid), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_col_widths(table, [col_w] * cols)
        first = images[0]["c"]
        overall = dict(first[0][2]).get("caption", "")
        for r, line in enumerate(grid):
            resolved = [(self._resolve(img["c"][2][0]), img) for img in line]
            heights = []
            for path, img in resolved:
                if path:
                    with PILImage.open(path) as im:
                        w_px, h_px = im.size
                    heights.append(min(usable * h_px / w_px, self.ROW_MAX_H))
            common = min(heights) if heights else 0
            for c, (path, img) in enumerate(resolved):
                cell = table.cell(r, c)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                if not path:
                    run = p.add_run(f"[Missing image: {img['c'][2][0]}]")
                    run.italic = True
                    run.font.color.rgb = RGBColor(0xC0, 0, 0)
                    continue
                width_attr = dict(img["c"][0][2]).get("width")
                if width_attr:
                    p.add_run().add_picture(str(path), width=Inches(min(self._img_width(path, img["c"][0][2]).inches, usable)))
                else:
                    with PILImage.open(path) as im:
                        aspect = im.size[0] / im.size[1]
                    p.add_run().add_picture(str(path), width=Inches(min(common * aspect, usable)))
                sub = _plain_text(img["c"][1]).strip()
                if sub:
                    sp = cell.add_paragraph()
                    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(4)
                    sp.paragraph_format.keep_with_next = number is not None
                    self._inlines(sp, img["c"][1], {})
                    for run in sp.runs:
                        run.font.size = Pt(self.s.body_size - 1)
        _no_borders(table)
        if number is not None:
            self._caption(container, "Figure", number, first[0][0], overall)
        else:
            spacer = container.add_paragraph()
            spacer.paragraph_format.space_after = Pt(4)

    def _caption(self, container, kind: str, number: int, label: str, inlines) -> None:
        """'Figure 3: text' with a live SEQ number and a bookmark that @fig:x cross-references point to."""
        cp = container.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = cp.paragraph_format
        pf.space_before, pf.space_after = (Pt(2), Pt(12)) if kind == "Figure" else (Pt(8), Pt(4))
        pf.keep_with_next = kind == "Table"
        size = self.s.body_size - 1
        bid = fields.bookmark_start(cp, fields.bookmark_name(label)) if label else None
        r = cp.add_run(f"{kind} ")
        r.bold, r.font.size = True, Pt(size)
        fields.add_field(cp, f"SEQ {kind} \\* ARABIC", str(number), bold=True, size_pt=size)
        if bid:
            fields.bookmark_end(cp, bid)
        text = _plain_text(inlines).strip() if not isinstance(inlines, str) else inlines.strip()
        if text:
            r = cp.add_run(": ")
            r.bold, r.font.size = True, Pt(size)
            before = len(cp.runs)
            if isinstance(inlines, str):
                cp.add_run(inlines.strip())
            else:
                self._inlines(cp, inlines, {})
            for run in cp.runs[before:]:
                run.font.size = Pt(size)
        text = re.sub(r"[{}\\$]", "", text).replace("_", "")
        self._captions[kind].append(f"{kind} {number}: {text}" if text else f"{kind} {number}")

    # ---- generated visuals (```plot, ```diagram, ```circuit)

    def _visual(self, container, block) -> None:
        import yaml

        from .diagrams import render_circuit, render_flowchart
        from .plots import make_plot

        attr, text = block["c"]
        kind = next(VISUAL_CLASSES[c.lower()] for c in attr[1] if c.lower() in VISUAL_CLASSES)
        kv = dict(attr[2])
        self._visual_n += 1
        out_dir = self.opt.figures_dir or Path(self._tmp.name)
        out = out_dir / f"figure_{self._visual_n:02d}_{kind}.png"
        try:
            if kind == "plot":
                spec = yaml.safe_load(text) or {}
                make_plot(spec, out, base_dir=self.assets_dir)
            elif kind == "diagram":
                render_flowchart(text, out)
            else:
                spec = yaml.safe_load(text)
                items = spec.get("elements", spec) if isinstance(spec, dict) else spec
                render_circuit(items, out)
        except Exception as exc:
            self.warnings.append(f"Could not draw {kind} #{self._visual_n}: {exc}")
            ep = container.add_paragraph()
            ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = ep.add_run(f"[{kind} could not be drawn: {exc}]")
            r.italic, r.font.color.rgb = True, RGBColor(0xC0, 0, 0)
            return
        default_w = {"plot": "85%", "diagram": None, "circuit": "70%"}[kind]
        size_kv = [["width", kv.get("width") or default_w]] if (kv.get("width") or default_w) else []
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        number = block.get("_fig_no")
        p.paragraph_format.keep_with_next = number is not None
        width = self._img_width(out, size_kv)
        p.add_run().add_picture(str(out), width=width)
        if number is not None:
            self._caption(container, "Figure", number, attr[0], block.get("_cap") or kv.get("caption", ""))
        self._record_figure(kind, attr, text, kv, out, width.inches, number)

    def _record_figure(self, kind, attr, text, kv, out: Path, width_in: float, number) -> None:
        from .figure_review import FigureRecord, critique, digest

        key = attr[0] or f"figure-{self._visual_n}"
        dg = digest(kind, text, kv)
        referenced = bool(attr[0]) and f"@{attr[0]}" in self.opt.body_markdown
        rec = FigureRecord(
            key=key, number=number, kind=kind, caption=kv.get("caption", ""), png=str(out), digest=dg,
            display_width_in=round(width_in, 2), approved=self.opt.approvals.get(key) == dg,
        )
        try:
            rec.issues = critique(kind, text, kv, out, width_in, referenced)
        except Exception as exc:
            rec.issues = [f"Automatic review failed: {exc}"]
        self.figures.append(rec)

    # ---- tables

    def _table(self, container, c, number=None) -> None:
        """Plain Word "Table Grid" by default: thin black lines, no shading, body font, bold header row.
        table_style="shaded" gives the navy-header / banded-row look instead."""
        _attr, caption, colspecs, head, bodies, foot = c
        rows = []
        for row in head[1]:
            rows.append((True, row[1]))
        for body in bodies:
            for row in body[2] + body[3]:
                rows.append((False, row[1]))
        for row in foot[1]:
            rows.append((False, row[1]))
        if not rows:
            return
        ncols = len(colspecs)

        cap_inlines = caption[1][0]["c"] if caption[1] else []
        if number is not None:
            self._caption(container, "Table", number, _attr[0], cap_inlines)

        table = container.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._table_borders(table)

        widths = self._column_widths(colspecs, rows, ncols)
        shaded = self.s.table_style == "shaded"
        body_i = 0
        for r_i, (is_head, cells) in enumerate(rows):
            band = None
            if is_head:
                band = TABLE_HEAD_FILL if shaded else None
                trpr = table.rows[r_i]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:tblHeader"))
            else:
                band = TABLE_BAND_FILL if shaded and body_i % 2 == 1 else None
                body_i += 1
            col = 0
            for cell in cells:
                _cattr, align, _rowspan, colspan, blocks = cell
                if col >= ncols:
                    break
                target = table.cell(r_i, col)
                if colspan > 1 and col + colspan - 1 < ncols:
                    target = target.merge(table.cell(r_i, col + colspan - 1))
                self._cell(target, blocks, is_head, align["t"], first_col=(col == 0), fill=band)
                col += colspan

        _set_col_widths(table, widths)

        spacer = container.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _column_widths(self, colspecs, rows, ncols) -> list[float]:
        total_w = CONTENT_W * 0.97
        given = [spec[1].get("c") if spec[1]["t"] == "ColWidth" else None for spec in colspecs]
        # Pandoc reports equal widths for long pipe tables (from the |---|---| row); that is not a real choice.
        if all(given) and max(given) - min(given) > 1e-6:
            return [total_w * w / sum(given) for w in given]
        # Size columns by how much text they hold, so a short label column stays narrow.
        lengths = [1.0] * ncols
        for _is_head, cells in rows:
            for i, cell in enumerate(cells[:ncols]):
                lengths[i] = max(lengths[i], len(_plain_text(cell[4])))
        weights = [min(max(n, 8), 60) ** 0.75 for n in lengths]
        return [total_w * w / sum(weights) for w in weights]

    def _table_borders(self, table) -> None:
        shaded = self.s.table_style == "shaded"
        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), TABLE_BORDER if shaded else "000000")
            borders.append(el)
        tblpr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        pads = (("top", 100), ("left", 115), ("bottom", 100), ("right", 115)) if shaded else \
            (("top", 29), ("left", 108), ("bottom", 29), ("right", 108))  # Word's own Table Grid padding
        for side, val in pads:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(val))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblpr.append(mar)

    def _cell(self, cell, blocks, is_head: bool, align: str, first_col: bool, fill: str | None) -> None:
        placeholder = cell.paragraphs[0]
        for b in blocks:
            self._block(cell, b, None)
        if len(cell.paragraphs) > 1:
            placeholder._p.getparent().remove(placeholder._p)
        shaded = self.s.table_style == "shaded"
        if fill:
            _shade(cell._tc.get_or_add_tcPr(), fill)
        if shaded and (is_head or first_col):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        font = self.s.table_font if shaded else self.s.body_font
        size = self.s.table_size if shaded else self.s.body_size
        for p in cell.paragraphs:
            pf = p.paragraph_format
            pf.space_after, pf.space_before, pf.line_spacing = Pt(2), Pt(2), 1.0 if not shaded else 1.15
            if shaded and (is_head or first_col):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = {
                    "AlignCenter": WD_ALIGN_PARAGRAPH.CENTER,
                    "AlignRight": WD_ALIGN_PARAGRAPH.RIGHT,
                }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
            for r in p.runs:
                r.font.name = font
                _set_fonts(r._r, font)
                r.font.size = Pt(size)
                if is_head:
                    r.bold = True
                    if shaded:
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # ---- code

    def _code(self, container, text: str) -> None:
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(10), 1.0
        pf.left_indent = pf.right_indent = Inches(0.05)
        _para_box(p)
        lines = text.expandtabs(4).split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = self.s.code_font
            _set_fonts(r._r, self.s.code_font)
            r.font.size = Pt(self.s.code_size)
            if i < len(lines) - 1:
                r.add_break()

    # ---- output

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        _normalise_order(self.doc.element)
        for st in self.doc.styles.element:
            _normalise_order(st)
        try:
            self.doc.save(path)
        except PermissionError:
            # Usually the previous version is still open in Word.
            alt = path.with_name(f"{path.stem} (new){path.suffix}")
            self.doc.save(alt)
            self.warnings.append(f"{path.name} is open in another program; saved as {alt.name} instead.")
            return alt
        finally:
            self._tmp.cleanup()
        return path


def build_report(
    *,
    style: TemplateStyle,
    profile: Profile,
    cover: CoverInfo,
    markdown: str,
    out_path: Path,
    assets_dir: Path | None = None,
    options: BuildOptions | None = None,
) -> BuildResult:
    b = ReportBuilder(style, profile, assets_dir, options)
    b.add_cover(cover)
    b.add_markdown(markdown)
    path = b.save(out_path)
    return BuildResult(path=path, warnings=b.warnings, figures=b.figures)


__all__ = ["CoverInfo", "BuildOptions", "BuildResult", "ReportBuilder", "build_report", "M"]
