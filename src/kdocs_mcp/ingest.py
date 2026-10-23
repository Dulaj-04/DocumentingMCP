"""Read source material: assignment briefs, lecture notes, earlier reports, spreadsheets."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".tex", ".c", ".h", ".py", ".m", ".ino", ".cpp", ".csv"}


@dataclass
class SourceText:
    path: Path
    pages: list[str]  # one entry per page (PDF) or a single entry for other formats

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)


def read_text(path: str | Path) -> SourceText:
    p = Path(path).expanduser()
    if not p.is_file():
        raise FileNotFoundError(f"File not found: {path}")
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return SourceText(p, pages)
    if ext == ".docx":
        import docx

        d = docx.Document(str(p))
        parts: list[str] = []
        body = d.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t"))
                if text.strip():
                    parts.append(text)
            elif tag == "tbl":
                for row in child.iter():
                    if row.tag.endswith("}tr"):
                        cells = []
                        for tc in row:
                            if tc.tag.endswith("}tc"):
                                cells.append(" ".join(t.text or "" for t in tc.iter() if t.tag.endswith("}t")).strip())
                        parts.append("| " + " | ".join(cells) + " |")
        return SourceText(p, ["\n".join(parts)])
    if ext in (".xlsx", ".xlsm"):
        return SourceText(p, [table_to_markdown(p, max_rows=200)])
    if ext in TEXT_EXT or not ext:
        return SourceText(p, [p.read_text(encoding="utf-8", errors="replace")])
    raise ValueError(f"Unsupported file type: {ext} (supported: PDF, DOCX, XLSX, CSV, TXT/MD and source code)")


# --------------------------------------------------------------------------- assignment briefs

_MODULE = re.compile(r"\b([A-Z]{2,3}\s?\d{4})\s*[-–—:]?\s*([A-Z][A-Za-z,&/ ]{3,80})?")
_DATE = re.compile(
    r"\b(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{1,2}(?:st|nd|rd|th)?\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?,?\s+\d{4})",
    re.I,
)
_QUESTION_START = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?(\d{1,2})\s*[.)]\s+(\S.*)$", re.I)
_SUBPART = re.compile(r"^\s*\(?([a-hivx]{1,4})[).]\s+(\S.*)$", re.I)
_MARKS = re.compile(r"[\[(]\s*(\d{1,3})\s*marks?\s*[\])]", re.I)


@dataclass
class Brief:
    module_code: str = ""
    module_name: str = ""
    title: str = ""
    due_date: str = ""
    questions: list[dict] = field(default_factory=list)
    instructions: list[str] = field(default_factory=list)


def parse_brief(text: str) -> Brief:
    b = Brief()
    lines = [ln.rstrip() for ln in text.splitlines()]

    m = _MODULE.search(text)
    if m:
        b.module_code = m.group(1).replace(" ", "")
        if m.group(2):
            b.module_name = m.group(2).strip().rstrip("-–— ").title()

    for ln in lines:
        if re.search(r"(due|deadline|submi(t|ssion))", ln, re.I):
            d = _DATE.search(ln)
            if d:
                b.due_date = d.group(1)
                break

    for ln in lines[:25]:
        s = ln.strip()
        if re.search(r"(assignment|lab(oratory)?\s*\d*|practical|report|tutorial|continuous assessment|CA\s*\d)", s, re.I) and 8 <= len(s) <= 120:
            if not _MODULE.fullmatch(s):
                b.title = s
                break

    current: dict | None = None
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        q = _QUESTION_START.match(s)
        sub = _SUBPART.match(s)
        if q and (current is None or int(q.group(1)) == current["number"] + 1 or int(q.group(1)) == 1):
            current = {"number": int(q.group(1)), "text": q.group(2).strip(), "parts": [], "marks": None}
            b.questions.append(current)
        elif current is not None and sub and len(sub.group(1)) <= 4:
            current["parts"].append(f"({sub.group(1)}) {sub.group(2).strip()}")
        elif current is not None:
            if current["parts"]:
                current["parts"][-1] += " " + s
            else:
                current["text"] += " " + s
        elif re.match(r"^[-•*·]\s+", s) or re.search(r"\b(must|should|submit|word limit|pages?|font|format)\b", s, re.I):
            b.instructions.append(s.lstrip("-•*· "))
        if current is not None:
            mk = _MARKS.search(s)
            if mk:
                current["marks"] = int(mk.group(1))
    for q in b.questions:
        q["text"] = _MARKS.sub("", q["text"]).strip()
    return b


# --------------------------------------------------------------------------- spreadsheets


def load_table(path: str | Path, sheet: str | None = None) -> list[list]:
    p = Path(path).expanduser()
    ext = p.suffix.lower()
    if ext == ".csv":
        with p.open(newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [row for row in csv.reader(f, dialect) if any(c.strip() for c in row)]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(p, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
        if rows:  # drop columns that are empty everywhere
            width = max(len(r) for r in rows)
            keep = [i for i in range(width) if any(i < len(r) and r[i] not in (None, "") for r in rows)]
            rows = [[r[i] if i < len(r) else None for i in keep] for r in rows]
        return rows
    raise ValueError("Spreadsheet must be .csv or .xlsx")


def _fmt(v, decimals: int | None) -> str:
    if v is None:
        return ""
    if isinstance(v, str) and decimals is not None:  # CSV cells arrive as text
        try:
            v = float(v)
        except ValueError:
            pass
    if isinstance(v, float):
        if v.is_integer() and decimals is not None:
            return str(int(v))  # whole numbers such as input codes stay whole
        if decimals is not None:
            return f"{v:.{decimals}f}"
        return f"{v:.6g}"
    return str(v).replace("|", "\\|").replace("\n", " ")


def table_to_markdown(
    path: str | Path,
    sheet: str | None = None,
    columns: list[str] | None = None,
    max_rows: int = 60,
    decimals: int | None = None,
) -> str:
    rows = load_table(path, sheet)
    if not rows:
        return ""
    header = [_fmt(c, None) or f"Column {i + 1}" for i, c in enumerate(rows[0])]
    body = rows[1:]
    if columns:
        idx = []
        for c in columns:
            if c not in header:
                raise ValueError(f"Column '{c}' not found. Available: {', '.join(header)}")
            idx.append(header.index(c))
        header = [header[i] for i in idx]
        body = [[r[i] if i < len(r) else None for i in idx] for r in body]
    truncated = len(body) > max_rows
    body = body[:max_rows]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for r in body:
        cells = [_fmt(r[i] if i < len(r) else None, decimals) for i in range(len(header))]
        lines.append("| " + " | ".join(cells) + " |")
    if truncated:
        lines.append(f"\n*({len(rows) - 1 - max_rows} more rows not shown)*")
    return "\n".join(lines)


def column_data(path: str | Path, column: str, sheet: str | None = None) -> list[float]:
    rows = load_table(path, sheet)
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if column not in header:
        raise ValueError(f"Column '{column}' not found in {Path(path).name}. Available: {', '.join(header)}")
    i = header.index(column)
    out = []
    for r in rows[1:]:
        v = r[i] if i < len(r) else None
        try:
            out.append(float(v))
        except (TypeError, ValueError):
            continue
    return out
