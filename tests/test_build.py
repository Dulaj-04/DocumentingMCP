import zipfile
from pathlib import Path

import anyio
import docx
import pytest
from PIL import Image

from documenting_mcp import server

EXAMPLES = Path(__file__).parent.parent / "examples"


@pytest.fixture(autouse=True)
def isolated_home(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCUMENTING_MCP_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("DOCUMENTING_MCP_REQUIRE_PLAN", "0")
    server.set_profile(name="A B C Student", reg_no="D/ENG/EE/25/0000", intake="42", output_dir=str(tmp_path / "out"))
    assets = tmp_path / "imgs"
    assets.mkdir()
    for name in ("debug_memory.png", "lcd_result.png", "dac_block_diagram.png"):
        Image.new("RGB", (800, 450), "white").save(assets / name)
    return assets


def _xml(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode()


def _build(template, example, assets, **kw):
    return server.create_document(
        template=template,
        title="Test Title",
        body_markdown=(EXAMPLES / example).read_text(encoding="utf-8"),
        module="ET2263 – Embedded Systems",
        submission_date="31/08/2026",
        assets_dir=str(assets),
        **kw,
    )


def test_lab_report(isolated_home):
    res = _build("lab_report", "lab_report.md", isolated_home)
    assert res["warnings"] == []
    d = docx.Document(res["docx"])
    texts = [p.text for p in d.paragraphs]
    assert "GENERAL SIR JOHN KOTELAWALA DEFENCE UNIVERSITY" in texts
    assert "ET2263 – EMBEDDED SYSTEMS" in texts
    assert any(t.startswith("Reg No\tD/ENG/EE/25/0000") for t in texts)
    headings = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert headings[:3] == ["Aim", "Apparatus", "Theory"]
    assert d.styles["Heading 1"].font.all_caps and d.styles["Heading 1"].font.underline
    xml = _xml(res["docx"])
    assert "<w:pgBorders" in xml
    assert "<m:oMathPara>" in xml  # display equation became a native Word equation
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "KDU UNIVERSITY"
    assert len(d.inline_shapes) == 3  # crest + 2 figures


def test_assignment_numbers_questions_but_not_references(isolated_home):
    res = _build("assignment", "assignment.md", isolated_home)
    d = docx.Document(res["docx"])
    questions = [p for p in d.paragraphs if p.style.name == "Heading 1"]
    assert len(questions) == 4
    assert all(q._p.pPr.numPr is not None for q in questions)
    assert questions[0].text.startswith("Define a Digital-to-Analog")
    refs = [p for p in d.paragraphs if p.text == "References"]
    assert refs and refs[0].style.name == "Heading 2"
    xml = _xml(res["docx"])
    assert xml.count("<m:oMath>") >= 6


def test_assignment_matches_master_format(isolated_home):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    res = _build("assignment", "assignment.md", isolated_home)
    d = docx.Document(res["docx"])
    sec = d.sections[0]
    assert sec.left_margin == Inches(0.5) and sec.right_margin == Inches(0.5)
    body = [p for p in d.paragraphs if p.text.startswith("A Digital-to-Analog Converter (DAC) is")][0]
    assert body.alignment == WD_ALIGN_PARAGRAPH.LEFT  # master is ragged-right
    assert body.paragraph_format.left_indent == Inches(0.375)  # aligned with question text
    bullet_heading = [p for p in d.paragraphs if p.text == "Binary Weighted Resistor DAC"][0]
    assert bullet_heading._p.pPr.numPr is not None  # "• Heading"
    xml = _xml(res["docx"])
    assert "w:fill=" not in xml  # plain Word grid tables: no colours
    assert 'w:color="000000"' in xml  # black grid lines
    header_run = d.tables[0].cell(0, 0).paragraphs[0].runs[0]
    assert header_run.bold and header_run.font.color.rgb is None
    assert header_run.font.name == "Times New Roman"


def test_shaded_tables_are_opt_in(isolated_home):
    res = _build("assignment", "assignment.md", isolated_home, table_style="shaded")
    xml = _xml(res["docx"])
    assert 'w:fill="2C3E7A"' in xml and 'w:fill="EDF0F8"' in xml


def test_missing_image_is_reported(isolated_home):
    res = server.create_document(
        template="ca_report", title="T", body_markdown="# Results\n\n![plot](nope.png)\n", assets_dir=str(isolated_home)
    )
    assert any("nope.png" in w for w in res["warnings"])


def test_question_prefix_is_stripped(isolated_home):
    res = server.create_document(template="assignment", title="T", body_markdown="# 1. What is a DAC?\n\nAnswer.\n")
    d = docx.Document(res["docx"])
    assert [p.text for p in d.paragraphs if p.style.name == "Heading 1"] == ["What is a DAC?"]


def test_tools_are_registered():
    tools = anyio.run(server.mcp.list_tools)
    names = {t.name for t in tools}
    assert {"get_profile", "set_profile", "get_writing_guide", "list_files", "create_document"} <= names
