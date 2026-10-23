from pathlib import Path

import docx
import pytest
from PIL import Image

from kdocs_mcp import server


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("KDOCS_MCP_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("KDOCS_MCP_REQUIRE_PLAN", "0")
    monkeypatch.setenv("KDOCS_MCP_IMAGES", str(tmp_path / "Downloads" / "KDocs Images"))
    server.set_profile(name="A B C Student", reg_no="D/ENG/EE/25/0000", intake="42", output_dir=str(tmp_path / "out"))
    return tmp_path


def _photos(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    out = []
    for name, size in (("IMG_0001.jpg", (1200, 800)), ("IMG_0002.JPG", (900, 900)), ("Screenshot (3).png", (800, 1200))):
        f = folder / name
        Image.new("RGB", size, "white").save(f)
        out.append(f)
    return out


def test_import_copies_and_renames_without_touching_originals(home):
    src = _photos(home / "phone")
    res = server.import_images([str(src[0].parent)], names=["Circuit Photo", "scope trace 1kHz", "block diagram"], project="ET2013 Lab 3")
    paths = [Path(i["path"]) for i in res["imported"]]
    assert [p.name for p in paths] == ["circuit-photo.jpg", "scope-trace-1khz.jpg", "block-diagram.png"]
    assert all(p.parent.name == "et2013-lab-3" for p in paths)
    assert all(f.exists() for f in src)  # originals untouched
    assert res["imported"][2]["shape"] == "portrait"
    listed = server.list_images("ET2013 Lab 3")["images"]
    assert len(listed) == 3


def test_wrong_number_of_names_is_ignored_with_a_note(home):
    src = _photos(home / "phone")
    res = server.import_images([str(src[0]), str(src[1])], names=["only-one"])
    assert [Path(i["path"]).name for i in res["imported"]] == ["img-0001.jpg", "img-0002.jpg"]
    assert res["problems"]


def test_duplicate_names_get_a_number(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0])], names=["setup"])
    res = server.import_images([str(src[1])], names=["setup"])
    assert Path(res["imported"][0]["path"]).name == "setup-2.jpg"


def test_rename_only_inside_the_library(home):
    src = _photos(home / "phone")
    assert "error" in server.rename_image(str(src[0]), "nope")
    assert src[0].exists()


def test_rename_updates_documents_and_images_resolve_from_the_library(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0])], names=["setup"])
    md = "# Q1\n\nSee @fig:s.\n\n![The bench setup](setup.jpg){#fig:s}\n"
    res = server.create_document(template="assignment", title="T", body_markdown=md, filename="r")
    assert not res["warnings"]  # found by bare name in the library
    out = server.rename_image("setup.jpg", "bench-setup", update_documents=[res["docx"]])
    assert out["renamed"] and out["documents"][0]["references"] == 1
    assert "bench-setup.jpg" in server.get_document_source(res["docx"])["body_markdown"]
    assert not out["documents"][0].get("warnings")


def test_images_on_one_line_become_an_aligned_row(home):
    src = _photos(home / "phone")
    server.import_images([str(f) for f in src], names=["a", "b", "c"])
    md = ('# Q1\n\nSee @fig:row.\n\n![(a) First](a.jpg){#fig:row caption="Three results"} ![(b) Second](b.jpg) '
          "![(c) Third](c.png)\n")
    res = server.create_document(template="assignment", title="T", body_markdown=md, filename="row")
    assert not res["warnings"]
    d = docx.Document(res["docx"])
    row = next(t for t in d.tables if len(t.columns) == 3)
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    extents = row._tbl.xpath(".//wp:extent")
    assert len(extents) == 3
    heights = [int(e.get("cy")) for e in extents]
    assert max(heights) - min(heights) < 400  # one common height (EMU)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Figure 1: Three results" in text or "Figure 1" in text
    assert "(a) First" in "\n".join(c.text for r in row.rows for c in r.cells)


# --------------------------------------------------------------------------- own folder, text under images, placement


def test_profile_image_folder_is_used(home):
    mine = home / "Downloads" / "My Lab Photos"
    mine.mkdir(parents=True)
    import os

    os.environ.pop("KDOCS_MCP_IMAGES", None)
    server.set_profile(image_dir=str(mine))
    src = _photos(home / "phone")
    res = server.import_images([str(src[0])], names=["bench"])
    assert Path(res["imported"][0]["path"]).parent == mine


def test_descriptions_become_text_under_the_image(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0])], names=["bench"], descriptions=["The lab bench with the circuit wired up"])
    listed = server.list_images()["images"][0]
    assert listed["description"].startswith("The lab bench")
    res = server.create_document(template="assignment", title="T", body_markdown="# Q1\n\nText.\n\n![The lab bench with the circuit](bench.jpg)\n", filename="cap")
    paras = [p.text for p in docx.Document(res["docx"]).paragraphs]
    assert "The lab bench with the circuit" in paras


def test_description_follows_a_rename(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0])], names=["a"], descriptions=["Output waveform"])
    server.rename_image("a.jpg", "output-wave")
    assert server.list_images()["images"][0]["description"] == "Output waveform"


BODY = """# Apparatus

We used a function generator, an oscilloscope and a breadboard.

# Procedure

The RC low-pass filter was wired on the breadboard as in the circuit diagram.

The input was set to 1 kHz and the output was read from the oscilloscope.

# Observations

The output amplitude at 1 kHz was lower than the input.
"""


def test_suggestions_match_images_to_the_right_sections(home):
    src = _photos(home / "phone")
    server.import_images([str(f) for f in src], names=["rc-filter-circuit", "scope-1khz-output", "holiday-beach"],
                         descriptions=["Circuit diagram of the RC low-pass filter wired on a breadboard",
                                       "Oscilloscope trace of the filter output at 1 kHz", "Friends at the beach"])
    out = server.suggest_image_placement(body_markdown=BODY)
    by_section = {s["section"]: [c["image"] for c in s["candidates"]] for s in out["suggestions"]}
    assert "rc-filter-circuit.jpg" in by_section["Procedure"]
    assert "scope-1khz-output.jpg" in by_section["Observations"]
    assert "holiday-beach.png" in out["not_matched_anywhere"]  # irrelevant images are left out
    after = next(c for s in out["suggestions"] if s["section"] == "Procedure" for c in s["candidates"] if c["image"].startswith("rc-"))
    assert "RC low-pass" in after["after_text"]


def test_place_image_inserts_after_the_chosen_paragraph(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0]), str(src[1])], names=["scope-1khz", "scope-2khz"],
                         descriptions=["Output trace at 1 kHz", "Output trace at 2 kHz"])
    res = server.create_document(template="assignment", title="T", body_markdown=BODY, filename="pl")
    out = server.place_image(res["docx"], "scope-1khz.jpg", "Procedure", after_text="input was set to 1 kHz", label="trace")
    assert "error" not in out and out["placed"]["section"] == "Procedure"
    body = server.get_document_source(res["docx"])["body_markdown"]
    assert body.index("read from the oscilloscope") < body.index("![Output trace at 1 kHz](scope-1khz.jpg){#fig:trace}") < body.index("# Observations")
    out = server.place_image(res["docx"], "scope-1khz.jpg", "Observations", side_by_side_with=["scope-2khz.jpg"], caption="Output at two frequencies")
    assert "error" not in out
    body = server.get_document_source(res["docx"])["body_markdown"]
    assert "![(a) Output trace at 1 kHz](scope-1khz.jpg){caption=\"Output at two frequencies\"} ![(b) Output trace at 2 kHz](scope-2khz.jpg)" in body
    assert "error" in server.place_image(res["docx"], "scope-1khz.jpg", "Procedure", after_text="no such words")


def test_checker_flags_an_image_with_no_text(home):
    src = _photos(home / "phone")
    server.import_images([str(src[0])], names=["bench"])
    res = server.create_document(template="assignment", title="T", body_markdown="# Q1\n\nText.\n\n![](bench.jpg)\n", filename="chk")
    report = server.check_document(res["docx"])
    assert any("no text under it" in i["message"] for i in report["issues"])
