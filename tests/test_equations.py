import zipfile

import docx
import pytest

from documenting_mcp import server
from documenting_mcp.latex_prep import prepare
from documenting_mcp.omml import M, latex_to_omml


@pytest.mark.parametrize(
    "tex, expected",
    [
        (r"V_{out} = V_{ref}", r"V_{\mathrm{out}} = V_{\mathrm{ref}}"),
        (r"V_{CE} + I_{in}", r"V_{\mathrm{CE}} + I_{\mathrm{in}}"),
        (r"b_{n-i} + x_{ij} + R_{k}", r"b_{n-i} + x_{ij} + R_{k}"),  # indices stay italic
        (r"V = 3.4375 V", r"V = 3.4375\,\text{V}"),
        (r"-5V", r"-5\,\text{V}"),
        (r"f = 50Hz", r"f = 50\,\text{Hz}"),
        (r"C = 10\mu F", r"C = 10\,\text{µF}"),
        (r"R = 2.2 k\Omega", r"R = 2.2\,\text{k}\Omega"),
        (r"2R + 4R", r"2R + 4R"),  # resistor multiples are not units
        (r"\SI{5}{\kilo\ohm}", r"5\,\mathrm{k\Omega}"),
        (r"\text{at 5 V}", r"\text{at 5 V}"),  # text is left alone
    ],
)
def test_scientific_normalisation(tex, expected):
    assert prepare(tex).latex == expected


def test_tag_and_environments():
    p = prepare(r"\begin{align} a &= b \label{eq:x} \tag{3} \end{align}")
    assert p.tag == "3"
    assert p.latex == r"\begin{aligned} a &= b   \end{aligned}"
    assert prepare(r"x = 1 \notag").notag


def test_omml_is_native_and_upright():
    (el,) = latex_to_omml([prepare(r"V_{out} = 5 V").latex])
    assert el.tag == f"{{{M}}}oMath"
    runs = {"".join(r.itertext()): r for r in el.iter(f"{{{M}}}r")}
    assert runs["out"].find(f"{{{M}}}rPr/{{{M}}}sty").get(f"{{{M}}}val") == "p"  # upright, one run
    unit = [r for t, r in runs.items() if t == "V" and r.find(f"{{{M}}}rPr/{{{M}}}nor") is not None]
    assert unit, "unit V should be a normal-text (upright) run"


def test_numbered_equations(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCUMENTING_MCP_HOME", str(tmp_path))
    monkeypatch.setenv("DOCUMENTING_MCP_REQUIRE_PLAN", "0")
    server.set_profile(output_dir=str(tmp_path / "out"))
    body = "# Q\n\n$$V = IR \\tag{7}$$\n\n$$P = VI$$\n\n$$E = Pt \\notag$$\n"
    res = server.create_document(template="assignment", title="T", body_markdown=body, number_equations=True)
    d = docx.Document(res["docx"])
    numbers = [t.cell(0, 2).text for t in d.tables]
    assert numbers == ["(7)", "(1)"]  # explicit tag, auto number, \notag skipped
    with zipfile.ZipFile(res["docx"]) as z:
        settings = z.read("word/settings.xml").decode()
    assert "m:mathPr" in settings and 'm:val="Cambria Math"' in settings
