import json
import zipfile
from pathlib import Path

import docx
import httpx
import pytest
from PIL import Image

from kdocs_mcp import ingest, project, references, server
from kdocs_mcp.plots import evaluate

EXAMPLES = Path(__file__).parent.parent / "examples"
REFS = [
    {"key": "razavi1995", "type": "book", "authors": ["Behzad Razavi"], "title": "Principles of Data Conversion System Design",
     "publisher": "IEEE Press", "address": "New York, NY, USA", "year": 1995},
    {"key": "sedra2015", "type": "book", "authors": ["Adel S. Sedra", "Kenneth C. Smith"], "title": "Microelectronic Circuits",
     "edition": 7, "publisher": "Oxford Univ. Press", "address": "New York, NY, USA", "year": 2015},
    {"key": "kester2005", "type": "book", "authors": ["Walt Kester"], "title": "The Data Conversion Handbook",
     "publisher": "Newnes", "address": "Burlington, MA, USA", "year": 2005},
]


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("KDOCS_MCP_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("KDOCS_MCP_REQUIRE_PLAN", "0")
    server.set_profile(name="A B C Student", reg_no="D/ENG/EE/25/0000", intake="42", output_dir=str(tmp_path / "out"))
    return tmp_path


def _xml(path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode()


def _full_demo(**kw):
    return server.create_document(
        template="assignment", title="DAC", body_markdown=(EXAMPLES / "full_assignment.md").read_text(encoding="utf-8"),
        module="ET2013 – Analog Electronics", references=REFS, filename="demo", **kw,
    )


# --------------------------------------------------------------------------- end goal guard


def test_default_output_is_the_master_format_with_page_numbers(home):
    res = server.create_document(template="assignment", title="T", body_markdown="# Q\n\nAnswer.\n")
    d = docx.Document(res["docx"])
    assert len(d.sections) == 2  # cover section + body section
    assert "PAGE" not in d.sections[0].footer._element.xml  # cover is never numbered
    assert "PAGE" in d.sections[1].footer._element.xml  # page numbers are on by default
    xml = _xml(res["docx"])
    assert "TOC" not in xml and "SEQ" not in xml
    with zipfile.ZipFile(res["docx"]) as z:
        assert b"updateFields" not in z.read("word/settings.xml")


def test_page_numbers_can_be_switched_off(home):
    res = server.create_document(template="assignment", title="T", body_markdown="# Q\n\nA.\n", page_numbers="none")
    assert len(docx.Document(res["docx"]).sections) == 1


# --------------------------------------------------------------------------- sources


def test_brief_parsing():
    b = ingest.parse_brief(
        "ET2013 - Analog Electronics\nAssignment on DACs\nSubmission deadline: 17/05/2026\n- Answer all questions.\n"
        "1. Define a DAC. [10 marks]\n2. Explain the R-2R ladder:\na) working principle\nb) advantages\n"
    )
    assert b.module_code == "ET2013" and b.due_date == "17/05/2026"
    assert [q["number"] for q in b.questions] == [1, 2]
    assert b.questions[0]["marks"] == 10
    assert b.questions[1]["parts"] == ["(a) working principle", "(b) advantages"]


def test_spreadsheet_to_table(home):
    csv = home / "r.csv"
    csv.write_text("Code,Vout (V)\n0,0\n1,0.3125\n", encoding="utf-8")
    md = server.spreadsheet_to_table(str(csv), decimals=2)["markdown"]
    assert md.splitlines()[0] == "| Code | Vout (V) |"
    assert "| 1 | 0.31 |" in md


def test_read_source_docx(home):
    res = server.create_document(template="assignment", title="T", body_markdown="# Q\n\nHello world.\n")
    assert "Hello world." in server.read_source(res["docx"])["text"]


# --------------------------------------------------------------------------- visuals


def test_visual_tools_produce_images(home):
    p = server.make_plot({"x_label": "t (s)", "series": [{"expr": "sin(2*pi*50*t)", "var": "t", "range": [0, 0.04]}]})
    d = server.make_diagram("flowchart LR\n A[In] --> B{Ok?}\n B -- Yes --> C((Out))\n B -- No --> A")
    c = server.make_circuit([{"type": "resistor", "dir": "right", "label": "R1"}, {"type": "capacitor", "dir": "down"}])
    for r in (p, d, c):
        assert "error" not in r, r
        with Image.open(r["png"]) as im:
            assert im.size[0] > 100


def test_formula_sandbox():
    with pytest.raises(ValueError):
        evaluate("__import__('os').system('echo hi')", {})
    with pytest.raises(ValueError):
        evaluate("x.__class__", {"x": 1})


# --------------------------------------------------------------------------- citations, captions, cross-references


def test_citations_captions_and_crossrefs(home):
    res = _full_demo()
    assert res["warnings"] == []
    d = docx.Document(res["docx"])
    text = "\n".join(p.text for p in d.paragraphs)
    assert "that number [1]." in text and "Figure 3 [2, p.\xa01012]" in text and "use it [3]." in text
    refs = [p.text for p in d.paragraphs if p.text.startswith("[")]
    assert refs[0].startswith("[1]\tB. Razavi, Principles of Data Conversion System Design.")
    assert refs[1].startswith("[2]\tA. S. Sedra and K. C. Smith, Microelectronic Circuits, 7th ed.")
    assert "Figure 1: Block diagram of an n-bit DAC" in text and "Table 1: Binary weighted" in text
    xml = _xml(res["docx"])
    assert "SEQ Figure" in xml and "SEQ Table" in xml and "SEQ Equation" in xml
    assert "REF _Ref_fig_block" in xml and 'w:name="_Ref_fig_block"' in xml
    assert "Eq. " in text  # @eq:dac
    figs = Path(res["figures_folder"])
    assert len(list(figs.glob("*.png"))) == 3  # diagram, plot, circuit


def test_missing_citation_and_crossref_are_reported(home):
    res = server.create_document(template="assignment", title="T", body_markdown="# Q\n\nSee @fig:nope and [@ghost].\n")
    assert any("@fig:nope" in w for w in res["warnings"])
    assert any("ghost" in w for w in res["warnings"])


def test_ieee_formats():
    article = {"type": "article", "authors": ["J. Doe", "Jane Roe", "Ann-Marie Lee"], "title": "A DAC", "journal": "IEEE JSSC",
               "volume": 40, "number": 3, "pages": "100–110", "year": 2005, "month": 3}
    assert references.to_plain(references.format_ieee(article)) == \
        'J. Doe, J. Roe, and A.-M. Lee, "A DAC," IEEE JSSC, vol. 40, no. 3, pp. 100–110, Mar. 2005.'
    web = {"type": "web", "organization": "Digilent", "title": "PmodCLS Reference Manual", "url": "https://x.y", "accessed": "2026-05-10"}
    assert references.to_plain(references.format_ieee(web)) == \
        'Digilent, "PmodCLS Reference Manual." [Online]. Available: https://x.y (accessed May 10, 2026)'


def test_doi_lookup_parses_crossref():
    def handler(request: httpx.Request) -> httpx.Response:
        assert "api.crossref.org/works/10.1109/5.771073" in str(request.url)
        return httpx.Response(200, json={"message": {
            "type": "journal-article", "title": ["Toward unique identifiers"], "author": [{"given": "Norman", "family": "Paskin"}],
            "container-title": ["Proceedings of the IEEE"], "volume": "87", "issue": "7", "page": "1208-1227",
            "issued": {"date-parts": [[1999, 7]]}, "DOI": "10.1109/5.771073"}})

    e = references.lookup("https://doi.org/10.1109/5.771073", transport=httpx.MockTransport(handler))
    assert references.to_plain(references.format_ieee(e)) == (
        'N. Paskin, "Toward unique identifiers," Proceedings of the IEEE, vol. 87, no. 7, pp. 1208–1227, Jul. 1999. '
        "doi: 10.1109/5.771073."
    )


def test_reference_library_round_trip(home):
    server.add_reference(key="st2017", type="datasheet", organization="STMicroelectronics", title="STM32F411xE", year=2017)
    assert [r["key"] for r in server.list_references()] == ["st2017"]
    res = server.create_document(template="assignment", title="T", body_markdown="# Q\n\nPer the datasheet [@st2017].\n")
    assert any(p.text.startswith("[1]\tSTMicroelectronics") for p in docx.Document(res["docx"]).paragraphs)


# --------------------------------------------------------------------------- opt-in furniture


def test_toc_page_numbers_and_header(home):
    res = _full_demo(toc=True, list_of_figures=True, page_numbers="bottom-center", header_text="ET2013")
    d = docx.Document(res["docx"])
    assert len(d.sections) == 2  # cover section has no number
    body = d.sections[1]
    assert body.header.paragraphs[0].text == "ET2013"
    assert "PAGE" in body.footer._element.xml
    xml = _xml(res["docx"])
    assert 'TOC \\o "1-3"' in xml and 'TOC \\h \\z \\c "Figure"' in xml
    with zipfile.ZipFile(res["docx"]) as z:
        assert b"updateFields" in z.read("word/settings.xml")


# --------------------------------------------------------------------------- editing


def test_section_edits_rebuild_in_place(home):
    res = _full_demo()
    path = res["docx"]
    src = server.get_document_source(path)
    assert [o["question"] for o in src["outline"] if o["level"] == 1] == [1, 2, None]  # References is not a question

    r = server.update_section(path, "Q2", "A new answer for question two [@kester2005].")
    assert r["docx"] == path and r["warnings"] == []
    r = server.insert_section(path, "# Define a DAC.\n\nA DAC converts codes to voltages.", before="Q1")
    r = server.replace_text(path, "codes to voltages", "digital codes to analog voltages")
    assert "1 occurrence" in r["change"]
    r = server.delete_section(path, "Block Descriptions")
    d = docx.Document(path)
    texts = [p.text for p in d.paragraphs]
    assert [p.text for p in d.paragraphs if p.style.name == "Heading 1"][:3] == [
        "Define a DAC.", "Explain the basic operation of a DAC using a suitable block diagram", "Explain the working principle of the Binary Weighted Resistor DAC"]
    assert "A DAC converts digital codes to analog voltages." in texts
    assert "Block Descriptions:" not in texts
    assert "A new answer for question two [2]." in texts  # citations renumber: razavi [1], kester [2]
    assert json.loads(Path(path).with_name("demo.kdocs.json").read_text())["template"] == "assignment"


def test_outline_targets():
    md = "# One\n\n## Working Principle\n\ntext\n\n# Two\n\n```\n# not a heading\n```\n\n# References\n"
    secs = project.sections(md)
    assert [s.title for s in secs] == ["One", "Working Principle", "Two", "References"]
    assert project.find_section(md, "Q2").title == "Two"
    assert project.find_section(md, "working").title == "Working Principle"
    new = project.insert_section(md, "# Three\n\nx")
    assert new.index("# Three") < new.index("# References")  # added before the reference list


# --------------------------------------------------------------------------- checks


def test_checker_finds_problems(home):
    body = ("# Q\n\nWe think the the output is 5V and this is importnt. The generalised colour of the centre is fine.\n\n"
            "*[Insert photo here]*\n\n# Empty\n\n# Q3\n\nText referencing nothing.\n\n![x](x.png){#fig:x}\n")
    res = server.create_document(template="assignment", title="T", body_markdown=body)
    rep = server.check_document(res["docx"], word_limit=10)
    msgs = " | ".join(i["message"] for i in rep["issues"])
    for expected in ("importnt", "the the", "'5 V'", "First-person", "Placeholder", "Section is empty", "Image not found", "word limit", "fig:x"):
        assert expected in msgs, expected
    assert "generalised" not in msgs and "colour" not in msgs and "centre" not in msgs  # British spelling is fine
    assert rep["summary"]["ready_to_submit"] is False


def test_similarity_flags_copied_paragraph(home):
    body = "# Q\n\nThe fundamental task of a DAC is to map each possible n-bit binary input code to a unique, proportionate analog output level in every system.\n\nAn original paragraph written entirely in the student's own words about converters and their many uses in practice today.\n"
    res = server.create_document(template="assignment", title="T", body_markdown=body)
    notes = home / "notes.txt"
    notes.write_text("Lecture 5: the fundamental task of a DAC is to map each possible n-bit binary input code to a unique, proportionate analog output level.")
    rep = server.check_similarity(res["docx"], [str(notes)])
    assert rep["paragraphs_flagged"] == 1 and rep["flagged"][0]["source"] == "notes.txt"


def test_checker_flags_ai_style(home):
    body = ("# Q\n\n**Type:** Group · **LOs:** LO1 · **Area:** Power\n\n"
            "It is worth noting that DACs play a crucial role in today's world. Furthermore, they seamlessly "
            "leverage robust ladders. Moreover, we delve into the landscape. Additionally, more text here.\n")
    res = server.create_document(template="assignment", title="T", body_markdown=body)
    msgs = " | ".join(i["message"] for i in server.check_document(res["docx"])["issues"])
    for expected in ("crucial", "seamless", "leverage", "robust", "delve", "landscape", "today's world",
                     "worth noting", "Furthermore/Moreover", "summary line"):
        assert expected in msgs, expected


# --------------------------------------------------------------------------- figure review and approval


FIG_BODY = """# Q

The output is in @fig:good.

```{.plot #fig:good caption="Output voltage against input code"}
x_label: Input code D
y_label: "Output voltage (V)"
series:
  - {label: ideal, x: [0, 1, 2, 3], y: [0, 1, 2, 3]}
```

```{.plot}
x_label: Time
series:
  - {x: [0, 1, 2], y: [1, 2, 3]}
  - {x: [0, 1, 2], y: [2, 3, 4]}
```

```{.diagram #fig:flow caption="Scan loop"}
flowchart TD
  A[Start] --> B{Device ready?}
  B --> C[Send]
  B --> A
  Z[Lonely block]
```
"""


def test_figures_need_student_approval(home):
    res = server.create_document(template="assignment", title="T", body_markdown=FIG_BODY)
    assert "action_required" in res and len(res["figures"]) == 3
    review = server.review_figures(res["docx"])
    assert review["all_approved"] is False and len(review["pending"]) == 3
    by_key = {f["key"]: f for f in review["figures"]}
    assert by_key["fig:good"]["issues"] == []  # labelled, captioned, referenced, units given
    bad_plot = " | ".join(by_key["figure-2"]["issues"])
    for expected in ("no unit", "Missing y label", "legend label", "No caption", "never referred"):
        assert expected in bad_plot, expected
    flow = " | ".join(by_key["fig:flow"]["issues"])
    assert "unlabelled branches" in flow and "Lonely block" in flow and "never referred" in flow
    assert Path(by_key["fig:good"]["png"]).is_file()

    # exporting and the checker are blocked until the student approves
    assert "not yet approved" in server.export_pdf(res["docx"])["error"]
    assert server.check_document(res["docx"])["summary"]["ready_to_submit"] is False

    out = server.approve_figures(res["docx"], ["Figure 1"])
    assert out["approved"] == ["Figure 1"]
    assert server.review_figures(res["docx"])["figures"][0]["approved"] is True

    # changing an approved figure sends it back for approval
    server.replace_text(res["docx"], "y: [0, 1, 2, 3]", "y: [0, 1.1, 2, 3]")
    assert server.review_figures(res["docx"])["figures"][0]["approved"] is False

    server.approve_figures(res["docx"], ["all"])
    assert server.review_figures(res["docx"])["all_approved"] is True


def test_bar_chart_keeps_category_labels(home):
    import matplotlib

    matplotlib.use("Agg")
    from kdocs_mcp.plots import make_plot
    import matplotlib.pyplot as plt

    seen = {}
    orig = plt.Figure.savefig

    def spy(self, *a, **k):
        seen["labels"] = [t.get_text() for t in self.axes[0].get_xticklabels()]
        return orig(self, *a, **k)

    plt.Figure.savefig = spy
    try:
        make_plot({"type": "bar", "categories": ["GS", "FD", "NR"], "y_label": "Iterations", "series": [{"y": [38, 9, 3]}]}, home / "b.png")
    finally:
        plt.Figure.savefig = orig
    assert seen["labels"] == ["GS", "FD", "NR"]


# --------------------------------------------------------------------------- regressions from the MA2263 trial


def test_numbered_references_heading_is_not_duplicated(home):
    server.add_reference(key="k1", type="book", authors=["Walt Kester"], title="Handbook", publisher="Newnes", year=2005)
    res = server.create_document(template="ca_report", title="T", body_markdown="# 1. Intro\n\nText [@k1].\n\n# 6. References\n")
    heads = [p.text for p in docx.Document(res["docx"]).paragraphs if p.style.name.startswith("Heading")]
    assert heads.count("6. References") == 1 and "References" not in heads


def test_author_suffix_and_preprint():
    assert references.format_authors(["A. Lozano III"]) == "A. Lozano III"
    assert references.format_authors(["Lozano, A., III"]) == "A. Lozano III"
    e = {"type": "preprint", "authors": ["S. E. Nwachukwu"], "title": "Power flow", "eprint": "arXiv:2506.23425", "year": 2025}
    assert references.to_plain(references.format_ieee(e)) == 'S. E. Nwachukwu, "Power flow," arXiv preprint arXiv:2506.23425, 2025.'


def test_absolute_value_bars_become_delimiters():
    from kdocs_mcp.latex_prep import prepare

    assert prepare(r"|V_i||V_j|").latex == r"\left|V_{i}\right|\left|V_{j}\right|" or \
        prepare(r"|V_i||V_j|").latex == r"\left|V_i\right|\left|V_j\right|"
    assert prepare(r"\|x\|").latex == r"\|x\|"


def test_checker_ignores_urls_plurals_and_reference_list(home):
    server.add_reference(key="w", type="web", organization="ST", title="Datasheet", url="https://www.st.com/ijeecs/xyzzq")
    body = ("# Q\n\nThe thermal capacitances and inductances were modelled; see https://www.st.com/ijeecs/xyzzq "
            "and doi: 10.11591/ijeecs.v34 for details [@w].\n\n# References\n")
    res = server.create_document(template="assignment", title="T", body_markdown=body)
    msgs = " | ".join(i["message"] for i in server.check_document(res["docx"])["issues"])
    for word in ("capacitances", "inductances", "https", "ijeecs", "xyzzq"):
        assert word not in msgs, word


def test_group_cover_and_department(home):
    members = [{"name": f"Student {i}", "reg_no": f"D/ENG/EE/25/000{i}"} for i in range(1, 6)]
    res = server.create_document(template="ca_report", title="Numerical Methods", body_markdown="# 1. Intro\n\nText.\n",
                                 department="Faculty of Engineering\nDepartment of Mathematics", members=members)
    d = docx.Document(res["docx"])
    texts = [p.text for p in d.paragraphs]
    assert "FACULTY OF ENGINEERING\nDEPARTMENT OF MATHEMATICS" in texts
    assert "Group Members" in texts and not any(t.startswith("Name\t") for t in texts)
    cover_table = d.tables[0]
    assert [c.text for c in cover_table.rows[0].cells] == ["No.", "Name", "Reg No"]
    assert cover_table.cell(5, 2).text == "D/ENG/EE/25/0005"


def test_first_hand_voice_checks(home):
    notes_style = ("# Observations\n\nThe notes record that the pump was started direct-on-line. According to the notes, "
                   "three pumps are installed. The photographed plate gives 1120 kW. This does not establish the starter. "
                   "It was not established which relay was used. The value cannot be verified. This should not be presented "
                   "as a measurement. The data does not prove anything.\n")
    res = server.create_document(template="lab_report", title="T", body_markdown=notes_style, filename="voice1")
    msgs = " ".join(i["message"] for i in server.check_document(res["docx"])["issues"])
    assert "instead of what was seen" in msgs and "Too many disclaimers" in msgs
    first_hand = ("# Observations\n\nThe pump was started direct-on-line. The plant engineers explained that three pumps are "
                  "installed. The nameplate showed a rated output of 1120 kW. The starter type could not be confirmed during the visit.\n")
    res = server.create_document(template="lab_report", title="T", body_markdown=first_hand, filename="voice2")
    msgs = " ".join(i["message"] for i in server.check_document(res["docx"])["issues"])
    assert "instead of what was seen" not in msgs and "Too many disclaimers" not in msgs


def test_writing_guide_has_first_hand_rules():
    guide = server.get_writing_guide("lab_report")
    assert "first-hand" in guide and "The nameplate showed" in guide and "Never invent an observation" in guide
